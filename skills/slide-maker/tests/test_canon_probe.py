#!/usr/bin/env python3
"""The canon rules that were made measurable must FIRE on the defect and stay silent on real decks.

This skill cites Duarte, Minto, CRAP, Mayer, Gestalt and cognitive load — all in prose, none
measured, which by its own enforcement invariant makes them advisory. `canon_probe.py` converts the
three that a built .pptx can actually decide. Each threshold was calibrated on 29 delivered decks /
349 slides, never on invented examples, and the calibration REJECTED two formulations outright:

  · Knaflic's declarative title, measured as overlap with the recorded takeaway: median 0.33, and
    the BEST titles score LOWEST ("Door one: buy" = 0.07) because a sharp title re-words on
    purpose. Judging "declarative" is taste; only the bare category label is mechanical.
  · Gestalt proximity: four formulations, all cry-wolf or silent — a .pptx records coordinates and
    no notion of which shapes belong together. See the module docstring.

So this file asserts BOTH directions for every rule that shipped: it catches the defect, and it
does not fire on a clean deck. A guard with 0 false positives that also never fires is not a guard.
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

import canon_probe as cp  # noqa: E402
import deckkit as dk      # noqa: E402

fails: list[str] = []
W, H = 13.333, 7.5
INK = dk.RGBColor.from_string("111111")
BODY = ("讨薪 The four gates each check one thing, and the build stops when any of them is red, "
        "so a broken page never reaches the reader at all.")


def check(cond, msg):
    if not cond:
        fails.append(msg)


def codes(found):
    return {c for c, _n, _m in found}


def add_chart(slide, *, labels, gridlines):
    """Build the chart with raw python-pptx on purpose.

    🔴 `deckkit.native_chart` has no `data_labels` switch at all, so the shape this rule targets
    CANNOT be produced by this skill's own helper — consistent with 0 of the 8 native charts in the
    corpus having it. The rule is a floor for decks that arrive from elsewhere (the redesign /
    critique path) or whose charts were hand-edited. Testing it through the helper would only prove
    the helper is clean.
    """
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    cd = CategoryChartData()
    cd.categories = ["a", "b", "c"]
    cd.add_series("v", (1, 2, 3))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.0), Inches(2.4),
                                Inches(5.0), Inches(3.0), cd)
    ch = gf.chart
    ch.plots[0].has_data_labels = labels
    ch.value_axis.has_major_gridlines = gridlines
    ch.value_axis.visible = gridlines
    return ch


def deck(td, title, body=BODY, notes="", chart=None):
    dk.set_palette(font="Helvetica Neue")
    prs = dk.blank_deck(W, H)
    s = dk.add_slide(prs)
    dk.text(s, 0.9, 0.6, 11.5, 0.9, [[(title, 30, INK, True, False)]])
    dk.text(s, 0.9, 2.4, 11.5, 2.0, [[(body, 18, INK, False, False)]])
    if notes:
        dk.speaker_notes(s, notes)
    if chart is not None:
        add_chart(s, **chart)
    out = td / "t.pptx"
    prs.save(str(out))
    return out


GATES = {"content": {"slides": [{"slide": 1, "role": "evidence", "takeaway": "x"}]}}

with tempfile.TemporaryDirectory() as _td:
    td = Path(_td)

    # ── clean baseline: nothing fires ────────────────────────────────────────────────────────────
    ok = deck(td, "The build stops at the first red gate",
              notes="Say that the point is not the count of gates but that one red one halts "
                    "everything downstream, and give the audience the shipping example.")
    check(codes(cp.check(ok, GATES)) == set(),
          "a clean deck produced findings: {}".format(cp.check(ok, GATES)))

    # ── 🔴 MAYER: notes that echo the slide ──────────────────────────────────────────────────────
    echo = deck(td, "The build stops at the first red gate", notes=BODY)
    check("NOTES ECHO SLIDE" in codes(cp.check(echo, GATES)),
          "verbatim notes did not fire — the rule that motivated the threshold is not enforced")
    # …and the legitimate case that sits highest in the real corpus stays silent. 0.53 was measured
    # across 253 slides; a threshold that flags elaboration would flag every well-noted deck.
    near = deck(td, "The build stops at the first red gate",
                notes="The four gates each check one thing. What the slide does not say is why the "
                      "order matters, and that is the part worth speaking to the room today.")
    check("NOTES ECHO SLIDE" not in codes(cp.check(near, GATES)),
          "notes that merely share vocabulary with the slide were flagged — real decks top out at "
          "0.53 similarity and would all fire")

    # ── 🔴 KNAFLIC: the bare category label ──────────────────────────────────────────────────────
    for bad in ("Overview", "背景", "Agenda", "  Summary:  ", "结论"):
        d = deck(td, bad)
        check("CATEGORY TITLE" in codes(cp.check(d, GATES)),
              "a bare category title {!r} passed — normalisation or the enumeration is broken"
              .format(bad))
    # a real declarative title, including one that shares almost nothing with its takeaway —
    # the case the overlap measure got WRONG and this one must not
    for good in ("Door one: buy", "Powered by wind, by default", "它不会把列表摘要当成已经查过",
                 "Overview of the three gates that block a release"):
        d = deck(td, good)
        check("CATEGORY TITLE" not in codes(cp.check(d, GATES)),
              "a real title {!r} was flagged as a category label".format(good))
    # 🔴 THE ARTIFACT-NAME CARVE, both directions. `Timeline`/`Roadmap` name the OBJECT on the
    # page, not the kind of page — a slide that IS one timeline is legitimately titled "Timeline",
    # and this gate BLOCKS, so a false positive costs a rename or a waiver. Found by running the
    # check on the first deck outside the 349-slide corpus (a layout fixture whose titles are form
    # names). The carve must not swallow the real labels next to it.
    for artifact in ("Timeline", "Roadmap", "时间线", "路线图"):
        check("CATEGORY TITLE" not in codes(cp.check(deck(td, artifact), GATES)),
              "{!r} was flagged — it names the artifact on the page, not a category of slide"
              .format(artifact))
    for still_bad in ("Overview", "Agenda", "Background", "背景", "目录"):
        check("CATEGORY TITLE" in codes(cp.check(deck(td, still_bad), GATES)),
              "{!r} stopped firing — the artifact-name carve widened into the real labels"
              .format(still_bad))

    # roles that legitimately carry a label title are exempt
    covergates = {"content": {"slides": [{"slide": 1, "role": "cover", "takeaway": "x"}]}}
    check("CATEGORY TITLE" not in codes(cp.check(deck(td, "Agenda"), covergates)),
          "a cover/section slide was charged for a label title")

    # ── 🔴 TUFTE: the same number printed twice ──────────────────────────────────────────────────
    twice = deck(td, "Three quarters, three numbers", chart={"labels": True, "gridlines": True})
    check("CHART SAYS IT TWICE" in codes(cp.check(twice, GATES)),
          "a chart with BOTH data labels and a value axis/gridlines passed — the Tufte rule is "
          "not enforced")
    # labels alone is the Tufte-APPROVED shape (erase the axis, keep the numbers): it must not fire
    once = deck(td, "Three quarters, three numbers", chart={"labels": True, "gridlines": False})
    check("CHART SAYS IT TWICE" not in codes(cp.check(once, GATES)),
          "direct-labelled bars with no axis were flagged — that is the shape Tufte asks FOR")
    # an axis with no labels is the corpus's own house style on all 8 charts: never flag it
    axis_only = deck(td, "Three quarters, three numbers", chart={"labels": False,
                                                                 "gridlines": True})
    check("CHART SAYS IT TWICE" not in codes(cp.check(axis_only, GATES)),
          "a plain axis-and-gridlines chart fired — all 8 charts in the real corpus are this "
          "shape and would every one of them be a false positive")

    # ── robustness: a deck with no notes, no chart, no gates record must not raise ───────────────
    bare = deck(td, "A title that says something")
    for g in (None, {}, {"content": {}}, {"content": {"slides": "not a list"}}):
        try:
            cp.check(bare, g)
        except Exception as exc:                                      # noqa: BLE001
            fails.append("check() raised on gates={!r}: {}".format(g, exc))

print("\n".join("FAIL " + f for f in fails if f) if [f for f in fails if f] else "", end="")
real = [f for f in fails if f]
print("[test_canon_probe] {}".format(
    "FAILED: {} problem(s)".format(len(real)) if real else "ok"))
sys.exit(1 if real else 0)
