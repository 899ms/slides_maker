#!/usr/bin/env python3
"""The two hand-off gates added for the template branch, both directions each.

`check_template_profile.py` — a registered template's profile.md must be OBEYED. The registry is
the one artefact in this skill with a memory across decks, and nothing read it. The important
property is not that it catches violations (it does) but that a profile with NO contract block
reports NOT CHECKED rather than clean: "there was nothing to check" and "everything checked out"
must never print the same sentence.

`check_fonts_resolve.py` — the faces a deck NAMES must resolve on the machine that MEASURED it.
Not portability (unknowable): every wrap/fit guard sits on `_measure_lines`, so a named-but-absent
face makes the build and the lint compute from the same wrong number and agree with each other.

Run: python3 tests/test_template_profile_and_fonts.py
"""
from __future__ import annotations

import json
import os
import sys
import tempfile
from pathlib import Path

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "scripts"))

from pptx import Presentation                                             # noqa: E402
from pptx.dml.color import RGBColor                                       # noqa: E402
from pptx.util import Inches, Pt                                          # noqa: E402

import check_fonts_resolve as CF                                          # noqa: E402
import check_template_profile as CT                                        # noqa: E402
import derive_template_contract as DTC                                       # noqa: E402

OK, BAD = [], []


def ck(cond, msg):
    (OK if cond else BAD).append(msg)
    print(("  ok   " if cond else "  ✗    ") + msg)


def _deck(path, *, face="Calibri", title_rgb="FFFFFF", cover_rect=None, n=1):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    for i in range(n):
        slide = prs.slides.add_slide(prs.slide_layouts[5])   # "Title Only" — a named layout
        t = slide.shapes.title
        t.text_frame.text = "A title"
        for r in t.text_frame.paragraphs[0].runs:
            r.font.size = Pt(28)
            r.font.name = face
            r.font.color.rgb = RGBColor.from_string(title_rgb)
        body = slide.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(6), Inches(1))
        run = body.text_frame.paragraphs[0].add_run()
        run.text = "Some body copy long enough to count as real text on the page."
        run.font.size = Pt(14)
        run.font.name = face
        if cover_rect and i == 0:
            from pptx.enum.shapes import MSO_SHAPE
            x, y, w, h = cover_rect
            # a FILLED shape — an empty textbox overlaps the rect and paints nothing, which is
            # exactly the false negative this check had until the fill test was added
            cov = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                         Inches(x), Inches(y), Inches(w), Inches(h))
            cov.fill.solid()
            cov.fill.fore_color.rgb = RGBColor.from_string("003C66")
    prs.save(path)
    layout_name = prs.slide_layouts[5].name
    return layout_name


def _profile(tmp, contract):
    d = Path(tmp) / "fake-template"
    d.mkdir(parents=True, exist_ok=True)
    body = "" if contract is None else (
        "\n## Machine-checkable contract\n```json\n%s\n```\n" % json.dumps(contract, indent=1))
    (d / "profile.md").write_text("# fake\n\nsome prose a human reads.\n" + body)
    return ("fake-template", d)


def _with_registry(entries, fn):
    import registry
    orig = registry.list_templates
    registry.list_templates = lambda: entries
    try:
        return fn()
    finally:
        registry.list_templates = orig


print("— template profile: a contract that is HONOURED stays silent")
with tempfile.TemporaryDirectory() as tmp:
    deck = os.path.join(tmp, "good.pptx")
    lay = _deck(deck)
    entry = _profile(tmp, {"match": {"layout_names": [lay], "slide_size_in": [10.0, 5.625]},
                           "layouts": {"content": lay},
                           "fonts": {"FONT": "Calibri"},
                           "title_color": "FFFFFF"})
    finds, facts = _with_registry([entry], lambda: CT.check(deck))
    ck(finds == [], "a deck matching its profile produces no findings (%r)" % (finds,))
    ck(set(facts["checked"]) == {"layouts", "fonts", "title_color"},
       "…and it reports exactly which keys it was able to check: %s" % facts["checked"])

print("\n— template profile: every declared key CATCHES its violation")
with tempfile.TemporaryDirectory() as tmp:
    deck = os.path.join(tmp, "bad.pptx")
    lay = _deck(deck, face="Helvetica", title_rgb="112233")
    entry = _profile(tmp, {"match": {"layout_names": [lay]},
                           "layouts": {"content": "A Layout That Is Not Used"},
                           "fonts": {"FONT": "Calibri"},
                           "title_color": "FFFFFF",
                           "must_cover": [{"rect": [7.14, 2.5, 2.54, 2.54], "why": "a grey block"}]})
    finds, _facts = _with_registry([entry], lambda: CT.check(deck))
    codes = {c for c, _ in finds}
    for want in ("LAYOUT OFF PROFILE", "FONT OFF PROFILE", "TITLE COLOUR OFF PROFILE",
                 "UNCOVERED TEMPLATE FURNITURE"):
        ck(want in codes, "%s fires on a deck that violates it" % want)

print("\n— template profile: must_cover is satisfied by a shape that actually covers it")
with tempfile.TemporaryDirectory() as tmp:
    deck = os.path.join(tmp, "covered.pptx")
    lay = _deck(deck, cover_rect=(7.14, 2.5, 2.54, 2.54))
    entry = _profile(tmp, {"match": {"layout_names": [lay]},
                           "must_cover": [{"rect": [7.14, 2.5, 2.54, 2.54], "why": "a grey block"}]})
    finds, _ = _with_registry([entry], lambda: CT.check(deck))
    ck(not finds, "covering the declared rect with a FILLED shape clears the finding (%r)" % (finds,))

with tempfile.TemporaryDirectory() as tmp:
    deck = os.path.join(tmp, "fakecover.pptx")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    sl = prs.slides.add_slide(prs.slide_layouts[5])
    sl.shapes.title.text_frame.text = "t"
    sl.shapes.add_textbox(Inches(7.14), Inches(2.5), Inches(2.54), Inches(2.54))   # paints nothing
    prs.save(deck)
    entry = _profile(tmp, {"match": {"layout_names": [prs.slide_layouts[5].name]},
                           "must_cover": [{"rect": [7.14, 2.5, 2.54, 2.54], "why": "a grey block"}]})
    finds, _ = _with_registry([entry], lambda: CT.check(deck))
    ck(any(c == "UNCOVERED TEMPLATE FURNITURE" for c, _ in finds),
       "an EMPTY TEXT BOX over the rect does not count as covering it")

print("\n— 🔴 a DESIGNED template has no layouts to fingerprint, and 10 of 11 registered ones are that")


def _painted_deck(path, colours, face="Helvetica Neue"):
    """A deck that paints exactly these colours — all a designed template leaves in the file."""
    from pptx.enum.shapes import MSO_SHAPE
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    for i, c in enumerate(colours):
        sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.3 + i * 0.4), Inches(0.3),
                                 Inches(0.35), Inches(0.35))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(c)
    tb = sl.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(6), Inches(1))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = "Body copy long enough to be the dominant face on this page."
    r.font.size = Pt(14)
    r.font.name = face
    prs.save(path)


DESIGNED = {"match": {"palette_any": ["14181F", "ECEFF4"]},
            "palette": {"core": ["14181F", "ECEFF4"],
                        "accents": ["34D1A6", "F2B04E"],
                        "expect": ["14181F", "1E242E", "ECEFF4", "34D1A6", "F2B04E"]},
            "fonts": {"FONT": "Helvetica Neue"}}

with tempfile.TemporaryDirectory() as tmp:
    deck = os.path.join(tmp, "d.pptx")
    _painted_deck(deck, ["14181F", "ECEFF4", "34D1A6"])
    finds, facts = _with_registry([_profile(tmp, DESIGNED)], lambda: CT.check(deck))
    ck(facts["template"] == "fake-template" and finds == [],
       "a deck painting the template's ground + ink BINDS on colour alone and passes — a designed "
       "template ships style.py and NO .pptx, so it has neither layout names nor a canvas size, "
       "and the layout fingerprint cannot describe it at all")
    ck(facts["palette"]["coverage"] == "3/5",
       "...and the REST of the palette is reported as coverage, never as a finding: measured, a "
       "4-slide deck built with modern-dark's own api painted 7 of its 11 declared colours, so a "
       "'60% of the palette' floor would have fired on a correct deck one slide shorter")

with tempfile.TemporaryDirectory() as tmp:
    deck = os.path.join(tmp, "d.pptx")
    _painted_deck(deck, ["14181F", "ECEFF4"])                  # ground + ink, not one accent
    finds, _f = _with_registry([_profile(tmp, DESIGNED)], lambda: CT.check(deck))
    ck(any(c == "PALETTE OFF PROFILE" for c, _m in finds),
       "a deck in the template's ground and ink with NONE of its accents is a finding — the "
       "accents are what make it this template rather than a grey box in its ground colour")

print("\n— 🔴 the RECORD binds where the fingerprint cannot: 'declared it, did not build it'")
_REC = {"interview": {"picks": [{"axis": "template", "source": "stated", "value": "fake-template"}]}}
with tempfile.TemporaryDirectory() as tmp:
    deck = os.path.join(tmp, "d.pptx")
    _painted_deck(deck, ["003C66", "FFFFFF"], face="Calibri")  # a stock deck, none of the palette
    ent = [_profile(tmp, DESIGNED)]
    try:
        _with_registry(ent, lambda: CT.check(deck))
        ck(False, "a stock deck binds to NOTHING by fingerprint")
    except RuntimeError as exc:
        ck("matches no registered template" in str(exc),
           "a stock deck binds to NOTHING by fingerprint — it has none of the template's colours, "
           "which is precisely the deck the gate most needs to catch")
    finds, facts = _with_registry(ent, lambda: CT.check(deck, gates=_REC))
    codes = sorted(c for c, _m in finds)
    ck(codes == ["FONT OFF PROFILE", "PALETTE OFF PROFILE"],
       "...and bound BY THE RECORDED NAME the same deck is caught twice: it declared the template "
       "and built in stock colours and the stock face. Only the record reaches this failure, which "
       "is why the gate reads it before falling back to the fingerprint")
    _pal_msg = next((m for c, m in finds if c == "PALETTE OFF PROFILE"), "")
    ck("#14181F" in _pal_msg and "ground and the text colour" in _pal_msg,
       "...and the palette finding is the CORE one, naming the ground and ink it paints neither "
       "of — the two branches share a code, so a test that reads only the code cannot tell the "
       "core check from the accent check and a mutant that deletes the core check survives")
    ck(facts["palette"]["core_missing"] == ["14181F", "ECEFF4"],
       "...with both missing core colours in the facts, where a report can quote them")

with tempfile.TemporaryDirectory() as tmp:
    deck = os.path.join(tmp, "d.pptx")
    _painted_deck(deck, ["14181F", "ECEFF4", "34D1A6"])
    bad_rec = {"interview": {"picks": [{"axis": "template", "source": "stated",
                                        "value": "a client deck.pptx"}]}}
    finds, facts = _with_registry([_profile(tmp, DESIGNED)],
                                  lambda: CT.check(deck, gates=bad_rec))
    ck(facts["template"] == "fake-template" and finds == [],
       "a recorded name that matches NO registered template falls back to the fingerprint rather "
       "than failing the deck — a user's own .pptx is a legitimate answer on that axis")
    ck(CT.recorded_template({"interview": {"picks": [{"axis": "template", "value": "<which one>"}]}})
       is None,
       "...and an unfilled scaffold placeholder is not read as a template name")

print("\n— the contract GENERATOR derives, and refuses to invent")
ck(DTC.from_style('BG = RGBColor(0x14, 0x18, 0x1F)\nFONT = "Helvetica Neue"\n')[0] == {"BG": "14181F"},
   "colours come out of the style module's own constants — a hand-copied hex is a number somebody "
   "remembered, which is what the never-invent floor is about")
_cols = {"BG": "14181F", "PANEL": "1E242E", "TEAL": "34D1A6"}
ck(DTC.identifying(_cols) == ["14181F", "34D1A6"],
   "the identifying pair is ground + LOUDEST colour by chroma, not by constant name — modern-dark "
   "calls its accent TEAL, nightdata calls its ORANGE, and a name-matching reader fingerprints "
   "both on two near-blacks")
ck(DTC._ink_for("14181F", ["14181F", "1E242E", "ECEFF4"]) == "ECEFF4",
   "the ink is the colour furthest from the ground in luminance — a text colour that does not "
   "contrast its ground is not a text colour")
with tempfile.TemporaryDirectory() as tmp:
    d = Path(tmp) / "t"
    d.mkdir()
    (d / "profile.md").write_text("# t\n\nprose\n\n## Machine-checkable contract\n```json\n{}\n```\n")
    ck(DTC.insert(d / "profile.md", DESIGNED) == "kept",
       "🔴 the generator REFUSES to overwrite a contract that is already there. lkeb-lumc's "
       "hand-written one carries layout names, a title colour and a must_cover rectangle that no "
       "palette scan can reproduce — and running --force over it once did exactly that")
    ck(DTC.insert(d / "profile.md", DESIGNED, force=True) == "wrote",
       "...and --force is the one way past it, so replacing a contract is always a decision")
_stock = DTC._deckkit_palette()
ck("003C66" in _stock and len(_stock) > 3,
   "deckkit's own palette is readable, which is what lets the generator refuse to build a "
   "fingerprint out of it")

print("\n— template profile: NO CONTRACT must be NOT CHECKED, never clean")
with tempfile.TemporaryDirectory() as tmp:
    deck = os.path.join(tmp, "x.pptx")
    _deck(deck)
    entry = _profile(tmp, None)
    # forced by name, because a contract-less profile has no `match` fingerprint to bind by
    try:
        _with_registry([entry], lambda: CT.check(deck, want="fake-template"))
        ck(False, "a profile with no contract block must RAISE, not return clean")
    except RuntimeError as exc:
        ck("Machine-checkable contract" in str(exc),
           "forcing a contract-less profile raises and says why: %s" % str(exc)[:58])
    # …and an UNforced run must name it as unmatchable rather than just shrugging
    try:
        _with_registry([entry], lambda: CT.check(deck))
        ck(False, "an unbindable deck must raise")
    except RuntimeError as exc:
        ck("declare no `## Machine-checkable contract`" in str(exc),
           "an unbound run names the contract-less templates as the likely reason")

print("\n— template profile: a deck matching NO profile binds to nothing")
with tempfile.TemporaryDirectory() as tmp:
    deck = os.path.join(tmp, "y.pptx")
    _deck(deck)
    entry = _profile(tmp, {"match": {"layout_names": ["Some Other Template's Layout"]},
                           "fonts": {"FONT": "Calibri"}})
    try:
        _with_registry([entry], lambda: CT.check(deck))
        ck(False, "a non-matching fingerprint must not bind")
    except RuntimeError as exc:
        ck("fingerprint" in str(exc), "a non-matching deck raises rather than borrowing a profile")

print("\n— fonts: an unresolved BODY face blocks, a one-character stray does not")
with tempfile.TemporaryDirectory() as tmp:
    deck = os.path.join(tmp, "f.pptx")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for txt, face in (("A heading nobody can measure", "Nonexistent Grotesk Pro"), ("x", "Ghost")):
        tb = s.shapes.add_textbox(Inches(0.6), Inches(1 + 1.5 * (txt == "x")), Inches(8), Inches(1))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text, r.font.name, r.font.size = txt, face, Pt(20)
    prs.save(deck)
    finds, facts = CF.check(deck)
    sev = {f: s_ for s_, f, _ in finds}
    ck(sev.get("Nonexistent Grotesk Pro") == "block",
       "a face carrying real text blocks (%r)" % sev.get("Nonexistent Grotesk Pro"))
    ck(sev.get("Ghost") == "note",
       "a face carrying one character is reported, not blocked (%r)" % sev.get("Ghost"))

print("\n— fonts: a deck whose faces ALL resolve produces no block")
# 🔴 The environment was never the thing under test. Earlier versions of this assertion built a
# deck in "a face that surely exists" — first a hardcoded {darwin: Helvetica, linux: DejaVu Sans},
# then whatever `an_installed_face()` reported — and BOTH failed on CI for reasons that had
# nothing to do with the code: the guessed face was absent, and then the discovered one still
# blocked for a face the assertion did not name, because the message printed no evidence. Three
# red runs bought one lesson: pin the LOGIC with a deterministic resolver, and let the
# environment-dependent half be the honest skip below.
with tempfile.TemporaryDirectory() as tmp:
    deck = os.path.join(tmp, "g.pptx")
    _deck(deck, face="Any Face At All")
    _real = CF._resolver
    CF._resolver = lambda: (lambda f: True)          # a host where everything resolves
    try:
        finds, _ = CF.check(deck)
    finally:
        CF._resolver = _real
    ck(finds == [],
       "when every face resolves, nothing is reported at all (got %r)"
       % [(s_, f) for s_, f, _ in finds])

# …and the same deck on a host where NOTHING resolves must block exactly the author-set face
with tempfile.TemporaryDirectory() as tmp:
    deck = os.path.join(tmp, "h.pptx")
    _deck(deck, face="Any Face At All")
    _real = CF._resolver
    CF._resolver = lambda: (lambda f: False)
    try:
        finds, _ = CF.check(deck)
    finally:
        CF._resolver = _real
    blocks = sorted(f for s_, f, _ in finds if s_ == "block")
    ck(blocks == ["Any Face At All"],
       "on a font-less host exactly the author-set face blocks; the theme's shipped default is a "
       "note (blocked: %r)" % blocks)

# the environment-dependent half, reported rather than asserted when the host cannot support it
_found = CF.an_installed_face()
if _found is None:
    print("  note: no resolvable face on this host — the live-environment direction is untestable "
          "here, reported rather than passed over")
else:
    with tempfile.TemporaryDirectory() as tmp:
        deck = os.path.join(tmp, "live.pptx")
        _deck(deck, face=_found)
        finds, _ = CF.check(deck)
        blocks = [(f, w[:70]) for s_, f, w in finds if s_ == "block"]
        ck(blocks == [],
           "a deck set in a face this host reports as installed (%s) produces no block — "
           "blocked instead: %r" % (_found, blocks))

print("\n— fonts: a SHIPPED DEFAULT that cannot resolve is a NOTE, not a block")
ck(CF.DEFAULT_FACES and "Calibri" in CF.DEFAULT_FACES,
   "the shipped-default set is derived from deckkit's source: %s" % sorted(CF.DEFAULT_FACES))
_real = CF._resolver
CF._resolver = lambda: (lambda face: False)          # a host with NO fonts at all — CI's situation
try:
    with tempfile.TemporaryDirectory() as tmp:
        deck = os.path.join(tmp, "nofonts.pptx")
        _deck(deck, face="Calibri")
        finds, _ = CF.check(deck)
        sev = {f: s_ for s_, f, _ in finds}
        ck(sev.get("Calibri") == "note",
           "on a font-less host, a deck in deckkit's DEFAULT face is a note (%r) — blocking it "
           "would refuse delivery on essentially every stock machine" % sev.get("Calibri"))
        ck([f for f in finds if f[0] == "block"] == [],
           "…and nothing in that deck blocks")
        deck2 = os.path.join(tmp, "chosen.pptx")
        _deck(deck2, face="Some Deliberately Chosen Face")
        finds2, _ = CF.check(deck2)
        ck(any(s_ == "block" and f == "Some Deliberately Chosen Face" for s_, f, _ in finds2),
           "…while a face the author CHOSE still blocks on the same host — the split is "
           "responsibility, not severity of the condition")
finally:
    CF._resolver = _real

print("\n— fonts: the derived default set matches deckkit's live module defaults")
import importlib, subprocess, json as _json
_out = subprocess.run([sys.executable, "-c",
    "import sys;sys.path.insert(0,'scripts');import deckkit,json;"
    "print(json.dumps([deckkit.FONT, deckkit.MONO, deckkit.EQFONT]))"],
    capture_output=True, text=True, cwd=os.path.dirname(HERE))
if _out.returncode == 0:
    live = set(_json.loads(_out.stdout))
    ck(live <= CF.DEFAULT_FACES,
       "every live deckkit default (%s) is in the derived set — if deckkit re-themes, this fails "
       "rather than silently blocking the new default" % sorted(live))
else:
    ck(False, "could not read deckkit's live defaults: %s" % _out.stderr[-120:])

print("\n— fonts: `ea` and `cs` typefaces are counted, not only `latin`")
ck("ea" in open(os.path.join(os.path.dirname(HERE), "scripts",
                             "check_fonts_resolve.py")).read(),
   "the CJK/complex-script slots are read — they carry the text on exactly the decks where "
   "substitution hurts most")

print()
print("%d passed, %d failed" % (len(OK), len(BAD)))
raise SystemExit(1 if BAD else 0)
