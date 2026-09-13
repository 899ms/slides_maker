#!/usr/bin/env python3
"""deckkit can draw a circle, and the library now exposes one.

🔴 MEASURED. A deck whose signature motif was a two-state RING shipped three rendered iterations
with a green SQUARE in its place. `box(round=True)` renders a rounded rectangle; `r=d/2` renders a
rounded rectangle; nothing raised, no lint code fired, and the author only found it by looking at
the render. deckkit used `MSO_SHAPE.OVAL` seventeen times internally — icon tiles, badges,
`concentric_rings`, node discs — and exposed it through no public helper, so an author reaching for
a circle found the rounded-rectangle switch and believed it.

Both halves are asserted: the new primitive really is an ellipse in the file, and `box(round=True)`
really is not — the second is what makes the first worth having.
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import deckkit as dk  # noqa: E402

fails: list[str] = []
GREEN = dk.RGBColor.from_string("3FBF6F")
PANEL = dk.RGBColor.from_string("1D2027")


def check(cond, msg):
    if not cond:
        fails.append(msg)


def prst_of(shape):
    """The OOXML preset geometry this shape actually carries."""
    import re
    m = re.search(r'<a:prstGeom[^>]*prst="([a-zA-Z0-9_]+)"', shape._element.xml)
    return m.group(1) if m else None


dk.set_palette(font="Helvetica Neue")
prs = dk.blank_deck(13.333, 7.5)
s = dk.add_slide(prs)

# ── 🔴 THE PRIMITIVE IS AN ELLIPSE, and the thing it replaces is not ─────────────────────────────
ring = dk.disc(s, 1.0, 1.0, 2.0, line=GREEN, line_w=8)
check(prst_of(ring) == "ellipse",
      "dk.disc did not produce an ellipse — it carries {!r}".format(prst_of(ring)))
for r in (None, 1.0, 0.5):
    kw = {"round": True} if r is None else {"round": True, "r": r}
    b = dk.box(s, 4.0, 1.0, 2.0, 2.0, line=GREEN, line_w=8, **kw)
    check(prst_of(b) != "ellipse",
          "box(round=True, r={!r}) now draws an ellipse — if that is deliberate, this whole "
          "helper is redundant and the docstring's history is wrong".format(r))

# ── fill / line grammar matches box ─────────────────────────────────────────────────────────────
dot = dk.disc(s, 7.0, 1.0, 0.5, fill=GREEN)
check("solidFill" in dot._element.xml, "a filled disc carries no solid fill")
both = dk.disc(s, 8.0, 1.0, 1.4, fill=PANEL, line=GREEN, line_w=3)
check("solidFill" in both._element.xml and "ln" in both._element.xml,
      "fill + line together did not both land")
hollow = dk.disc(s, 10.0, 1.0, 1.0, line=GREEN)
check("noFill" in hollow._element.xml, "a disc with no fill is not hollow")

# ── an ellipse when h differs ───────────────────────────────────────────────────────────────────
from pptx.util import Emu  # noqa: E402
el = dk.disc(s, 1.0, 4.0, 2.0, h=1.0, line=GREEN)
check(abs(Emu(el.width).inches - 2.0) < 1e-6 and abs(Emu(el.height).inches - 1.0) < 1e-6,
      "h= did not produce an ellipse: {:.2f} x {:.2f}".format(Emu(el.width).inches,
                                                              Emu(el.height).inches))

# ── 🔴 NO INHERITED THEME SHADOW — the second half of the same trap ──────────────────────────────
# A shape added with slide.shapes.add_shape outside deckkit keeps the theme <p:style>, and
# LibreOffice draws a soft drop shadow under it (INHERITED_EFFECT). Its only cure was the PRIVATE
# `_flat`, which an author has to read the source to find.
check("<p:style>" not in ring._element.xml,
      "dk.disc leaves the theme <p:style> on the shape — it must be _flat()ed like box")
raw = s.shapes.add_shape(__import__("pptx.enum.shapes", fromlist=["MSO_SHAPE"]).MSO_SHAPE.OVAL,
                         Emu(0), Emu(0), Emu(914400), Emu(914400))
check("<p:style>" in raw._element.xml,
      "a raw add_shape no longer carries <p:style>; if python-pptx changed, this helper's "
      "shadow-flattening reason needs re-checking rather than assuming")

# ── it is placed like box: (x, y) is the TOP-LEFT, not the centre ───────────────────────────────
placed = dk.disc(s, 3.0, 5.0, 1.0, fill=GREEN)
check(abs(Emu(placed.left).inches - 3.0) < 1e-6 and abs(Emu(placed.top).inches - 5.0) < 1e-6,
      "disc treats (x, y) as a centre; every other primitive takes the top-left and a disc that "
      "differs cannot drop into a columns()/rows() rect")

# ── it composes with the motif / accessibility machinery ────────────────────────────────────────
try:
    dk.tag_motif(dk.disc(s, 5.0, 5.0, 0.6, fill=GREEN), loud=False)
except Exception as exc:                                              # noqa: BLE001
    fails.append("tag_motif does not accept a disc: {}".format(exc))

# ── degenerate input fails loudly rather than drawing nothing ───────────────────────────────────
for bad in ({"d": 0}, {"d": -1}, {"d": 1.0, "h": 0}):
    try:
        dk.disc(s, 0.5, 0.5, **bad)
        fails.append("disc({!r}) drew something instead of raising".format(bad))
    except ValueError:
        pass

# ── and the deck still saves + lints ────────────────────────────────────────────────────────────
with tempfile.TemporaryDirectory() as td:
    try:
        dk.lint_layout(prs, strict=False)
        prs.save(str(Path(td) / "t.pptx"))
    except Exception as exc:                                          # noqa: BLE001
        fails.append("a deck containing discs failed to lint/save: {}".format(exc))

print("\n".join("FAIL " + f for f in fails) if fails else "", end="")
print("[test_disc_primitive] {}".format(
    "FAILED: {} problem(s)".format(len(fails)) if fails else "ok"))
sys.exit(1 if fails else 0)
