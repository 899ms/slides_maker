#!/usr/bin/env python3
"""The cheap build-time loop must not report clean on a fault the expensive render-time loop finds.

🔴 MEASURED. `lint_layout`'s footer check asked whether a CARD reaches the actual footer chrome
row; `lint_deck`'s asks whether any TEXT's ink dips into the reserved band. Different questions —
so a low text block passed the build-time loop BY CONSTRUCTION and failed at render. On one real
build that cost FIVE consecutive build+render rounds, three of them refused by the LOOP BREAKER
for nudging constants at a fault the fast loop could not see.

And the two ink models disagree in the direction that hurts: `lint_deck._rbox` reads 0.020-0.040in
LOWER than `deckkit._ink_rect`, and the render-time check adds a further 0.04 pad. An author who
designs to `_ink_rect` — which this skill explicitly tells them to do — lands inside the
render-time failure band.

THE PROPERTY, which is what this file pins: on any deck, every slide the RENDER-time check flags
is also flagged by the BUILD-time one. Firing earlier is the point (`FOOTER_BAND_PAD`); firing
later is the bug. Verified across the delivered decks on this machine and on a built fixture.

🔴 The footer-identification rule is lint_deck's own, READ from its source rather than
approximated: text whose FRAME TOP sits below `sh - 0.6`. Two approximations were tried first — an
absolute 0.14in from the bottom edge, then 4% of canvas height — and both missed the same real
intrusion, because they were tuned against a different ink model. Tuning a threshold toward a rule
you have not read is the loop the LOOP BREAKER exists to stop, and it applies to whoever writes
the checker too.
"""
from __future__ import annotations

import re
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

import deckkit as dk  # noqa: E402

fails: list[str] = []


def check(cond, msg):
    if not cond:
        fails.append(msg)


def render_flags(pptx: Path) -> set[int]:
    r = subprocess.run([sys.executable, str(SCRIPTS / "lint_deck.py"), str(pptx)],
                       capture_output=True, text=True)
    return {int(m) for m in re.findall(r"slide (\d+): FOOTER-ZONE", r.stdout + r.stderr)}


def build_flags(prs) -> set[int]:
    return {n for n, _lvl, _code, _msg in dk._footer_band_faults(prs)}


dk.set_palette(font="Helvetica Neue")
INK = dk.RGBColor.from_string("111111")
W, H = 10.0, 5.625

# ── 🔴 THE SUPERSET PROPERTY, ON A REAL DELIVERED DECK ───────────────────────────────────────────
# A synthetic fixture cannot exercise this: the render-time limit is keyed to where the FOOTER
# actually sits (`min(footer tops) - 0.08`), so a fixture whose footer hugs the bottom edge has a
# limit no synthetic body line reaches, and the check never fires. Tried at three body heights —
# the render-time check stayed silent at all of them while the build-time one spoke, which is the
# safe direction but proves nothing about the unsafe one. The decks that DO exercise it are real
# ones; this repo has a delivered deck carrying three unresolved FOOTER-ZONE intrusions, and that
# is the test data.
from pptx import Presentation  # noqa: E402

CORPUS = Path.home() / "Downloads"
real = []
for rel in ("slides_skill_test/melbourne-first-trip/melbourne-first-trip.pptx",
            "agenthalo-intro/agenthalo-intro.pptx",
            "job-hunt-intro-zh/job-hunt-intro-zh.pptx"):
    p_ = CORPUS / rel
    if p_.exists():
        real.append(p_)

if not real:
    print("  (no delivered decks readable here — macOS can withhold ~/Downloads mid-session; the "
          "synthetic half below still runs)")
for p_ in real:
    try:
        prs_ = Presentation(str(p_))
    except Exception as exc:                                          # noqa: BLE001
        fails.append("could not open {}: {}".format(p_.name, exc))
        continue
    rf, bf = render_flags(p_), build_flags(prs_)
    missed = sorted(rf - bf)
    check(not missed,
          "{}: slide(s) {} are flagged at RENDER time and NOT at build time — the cheap loop "
          "reports clean on a fault the expensive loop finds, which is the whole defect"
          .format(p_.name, missed))

# ── a deck built through content_band() trips NEITHER ───────────────────────────────────────────
with tempfile.TemporaryDirectory() as _td:
    td = Path(_td)
    prs2 = dk.blank_deck(W, H)
    for _ in range(3):
        s = dk.add_slide(prs2)
        dk.text(s, 0.6, 0.4, 8.8, 0.6, [[("A title that says something", 24, INK, True, False)]])
        x, y, w, h = dk.content_band(s, top=1.4)
        dk.text(s, x, y, w, 1.0, [[("body placed through content_band", 14, INK, False, False)]])
        dk.text(s, 0.6, H - 0.42, 6.0, 0.3, [[("Source: x", 10, INK, False, False)]])
    deck2 = td / "clean.pptx"
    prs2.save(str(deck2))
    check(not build_flags(prs2),
          "a deck whose body came from content_band() was flagged — the check punishes the very "
          "helper the fix tells authors to use")
    check(not render_flags(deck2),
          "the clean fixture trips the render-time check too; it is not a valid negative control")

    # …and text pushed deliberately low DOES trip the build-time check, so it is not inert
    prs3 = dk.blank_deck(W, H)
    s = dk.add_slide(prs3)
    dk.text(s, 0.6, 0.4, 8.8, 0.6, [[("A title", 24, INK, True, False)]])
    dk.text(s, 0.6, H - 0.62, 6.0, 0.5, [[("a line shoved into the band", 14, INK, False, False)]])
    dk.text(s, 0.6, H - 0.42, 6.0, 0.3, [[("Source: x", 10, INK, False, False)]])
    check(build_flags(prs3),
          "text pushed into the reserved band was not flagged at build time — the check is inert")

# ── the pad exists and points the safe way ──────────────────────────────────────────────────────
check(getattr(dk, "FOOTER_BAND_PAD", 0) > 0,
      "FOOTER_BAND_PAD is gone or non-positive — without it this check fires at the same line as "
      "the render-time one and clearing it here no longer guarantees clearing it there")

print("\n".join("FAIL " + f for f in fails) if fails else "", end="")
print("[test_footer_band_superset] {}".format(
    "FAILED: {} problem(s)".format(len(fails)) if fails else "ok"))
sys.exit(1 if fails else 0)
