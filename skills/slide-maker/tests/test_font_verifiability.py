#!/usr/bin/env python3
"""A font inside an app bundle is a font this pipeline cannot verify.

Portability was only ever framed as "will the RECIPIENT have it". The failure that actually cost a
build was the other one: on macOS, Microsoft Office keeps Calibri / Cambria / Aptos INSIDE its own
app folder rather than installing them system-wide. PowerPoint renders them; LibreOffice — the
render loop — and the width-measurement path cannot see them at all. So a deck set in Calibri on
such a machine is laid out against a substitute's metrics (voiding `measure_text`, `vstack`,
`bottom_callout`, `fit_text_size`) and then "verified" against a render of a face nobody will see,
with both linters green the whole way.

Item 10 used to answer that case with "may be fine where the deck is presented" — which is the
wrong risk, stated reassuringly. These hold the distinction.
"""
import subprocess
import sys
import tempfile
import warnings
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))
warnings.simplefilter("ignore")

import preflight_check as pf                # noqa: E402
import deckkit as dk                        # noqa: E402

ok, bad = [], []


def check(cond, why):
    (ok if cond else bad).append(why)


# ------------------------------------------------------------------ the family-name matcher
check(pf._norm_face("Times New Roman") == "timesnewroman",
      "a family name normalises to alphanumerics, so 'Times New Roman' matches 'Times New Roman.ttf'")
check(pf._norm_face("") == "" and pf._bundled_only("") is None,
      "an empty face name is not matched against every file in every bundle")

# 🔴 The direction of the prefix test is the whole correctness of this check.
_stem, _want = pf._norm_face("Cambria"), pf._norm_face("Cambria Math")
check(not _stem.startswith(_want),
      "🔴 `Cambria.ttc` must NOT answer for `Cambria Math` — they are different fonts, and a "
      "machine with Office routinely carries the first and not the second (measured). The test is "
      "stem.startswith(family), never the reverse, or every math deck would be told its math font "
      "is present when the formulas are about to tofu")
check(pf._norm_face("Calibrib").startswith(pf._norm_face("Calibri")),
      "...while a weight file (`Calibrib.ttf`) DOES answer for its family, which is what makes the "
      "bundle scan find Calibri at all")
check(pf._bundled_only("NoSuchFaceAnywhere12345") is None,
      "a face in no bundle returns None rather than a guess")


# ------------------------------------------------------------------ real decks, never fakes
# 🔴 These arms used to drive item 10 with fake objects exposing only `run.font.name`. That fake IS
# the defect: python-pptx's `font.name` is the <a:latin> slot alone, so a fake built from it could
# never contain a Chinese face, and item 10 passed a Chinese deck whose CJK face resolved nowhere
# (measured: 24 glyphs laid out at 60% of their true width, "all 1 font(s) resolve locally:
# ['Calibri']"). Every deck below is a real .pptx, saved and reopened.
import io                                                                  # noqa: E402
import os                                                                  # noqa: E402
from lxml import etree                                                     # noqa: E402
from pptx import Presentation                                              # noqa: E402
from pptx.oxml.ns import qn                                                # noqa: E402
from pptx.util import Inches                                               # noqa: E402
import check_fonts_resolve as cfr                                          # noqa: E402


def _reopen(prs):
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return Presentation(buf)


def _slot(rPr, slot, face):
    el = rPr.find(qn("a:" + slot))
    if el is None:
        el = rPr.makeelement(qn("a:" + slot), {})
        rPr.append(el)                              # fixtures set slots in schema order
    el.set("typeface", face)


def _run(paragraph, text, faces, lang=None):
    r = paragraph.add_run()
    r.text = text
    rPr = r._r.get_or_add_rPr()
    for slot in ("latin", "ea", "cs"):
        if slot in faces:
            _slot(rPr, slot, faces[slot])
    if lang:
        rPr.set("lang", lang)
    return r


def _deck(rows, theme=None, layout=6):
    """rows = [(text, {slot: face}, lang)] -> a reopened Presentation."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[layout])
    for i, (text, faces, lang) in enumerate(rows):
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.4 + 0.7 * i), Inches(8), Inches(0.6))
        _run(tb.text_frame.paragraphs[0], text, faces, lang)
    if theme:
        _patch_theme(prs, theme)
    return _reopen(prs)


def _patch_theme(prs, theme):
    """theme = {("minor"|"major", "ea"|"cs"|"script:Hans"): face}; empty script list otherwise."""
    for rel in prs.slide_masters[0].part.rels.values():
        if not rel.reltype.endswith("/theme"):
            continue
        part = rel.target_part
        root = etree.fromstring(part.blob)
        for scheme, tag in (("major", "majorFont"), ("minor", "minorFont")):
            el = root.find(".//" + qn("a:" + tag))
            for node in el.findall(qn("a:font")):
                el.remove(node)                     # deterministic: no inherited script faces
            for (sch, slot), face in theme.items():
                if sch != scheme:
                    continue
                if slot.startswith("script:"):
                    node = el.makeelement(qn("a:font"), {"script": slot[7:], "typeface": face})
                    el.append(node)
                else:
                    el.find(qn("a:" + slot)).set("typeface", face)
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


class _Resolver(object):
    """A host on which exactly `unresolved` fail to resolve — the LOGIC, pinned apart from fonts."""

    def __init__(self, unresolved):
        self.unresolved, self.real = set(unresolved), cfr._resolver

    def __enter__(self):
        cfr._resolver = lambda: (lambda face: face not in self.unresolved)

    def __exit__(self, *exc):
        cfr._resolver = self.real


# ------------------------------------------------------- the two failures are reported apart
_live = cfr._resolver()
_nowhere = _deck([("A headline that is long enough", {"latin": "ZzzNotARealFace"}, None),
                  ("and a second line of body text", {"latin": "AlsoNotReal"}, None)])
if _live is not None:
    check(_live("ZzzNotARealFace") is False,
          "premise: a face named nowhere really does not resolve on this host")
status, msg = pf.item10_fonts(_nowhere)
check(status == "ADVISORY" and "presented" in msg and "stand-in" in msg,
      "a face installed NOWHERE stays advisory — the risk really is partly the presenter's machine, "
      "and CI not having Helvetica Neue must not fail a deck — but it no longer only REASSURES: it "
      "says the geometry was measured in a stand-in ({}: {})".format(status, msg[:90]))


def _bundle_faces():
    """{"only": (family, app) bundled and NOT resolvable here, "both": family bundled AND resolvable}.

    🔴 ASK the machine. This arm hardcoded Calibri as "the bundle-only face", which was true until
    the user followed the skill's own advice and copied Calibri into ~/Library/Fonts — after which
    the premise was false, the product code was right, and the suite stayed red on the one machine
    where the arm runs at all (CI has no Office bundle, so it skips there).
    """
    out = {}
    ok = cfr._resolver()
    try:
        from matplotlib import ft2font
    except Exception:                                                  # noqa: BLE001
        return out
    if ok is None:
        return out
    for app, d in pf._BUNDLED_FONT_DIRS:
        try:
            names = sorted(os.listdir(d))
        except OSError:
            continue
        for fn in names:
            if os.path.splitext(fn)[1].lower() not in (".ttf", ".otf", ".ttc"):
                continue
            try:
                fam = ft2font.FT2Font(os.path.join(d, fn)).family_name
            except Exception:                                          # noqa: BLE001
                continue
            if not fam or not pf._bundled_only(fam):
                continue
            verdict = ok(fam)
            if verdict is False and "only" not in out:
                out["only"] = (fam, app)
            elif verdict is True and "both" not in out:
                out["both"] = fam
            if len(out) == 2:
                return out
    return out


_bf = _bundle_faces()
if "only" in _bf:
    _fam, _app = _bf["only"]
    status, msg = pf.item10_fonts(_deck([("A headline set in a bundled face", {"latin": _fam}, None)]))
    check(status == "ADVISORY" and "application bundle" in msg and "SUBSTITUTE" in msg,
          "🔴 a face present ONLY inside an app bundle is reported as its own class — the message "
          "says the GEOMETRY was measured against a substitute and the render verified the wrong "
          "face, not the reassuring 'may be fine where the deck is presented' ({} inside {})"
          .format(_fam, _app))
    check("~/Library/Fonts" in msg and "Times New Roman" in msg,
          "...and it names both fixes: a system-wide face, or making the bundled file visible")
    status2, msg2 = pf.item10_fonts(_deck([("A headline set in a bundled face", {"latin": _fam}, None),
                                           ("and a line in a face found nowhere", {"latin": "ZzzNotARealFace"}, None)]))
    check("application bundle" in msg2 and "NOWHERE" in msg2,
          "...and when BOTH failures are present, both are reported — the bundled one must not "
          "swallow a genuinely missing math font, which is the tofu that reaches the audience")
else:
    check(True, "no face on this machine lives ONLY inside an app bundle — the bundle-detection arm "
                "is skipped, not faked (it runs wherever Office keeps a face the system lacks)")
if "both" in _bf:
    status, msg = pf.item10_fonts(_deck([("A headline set in a face installed twice", {"latin": _bf["both"]}, None)]))
    check(status == "PASS",
          "🔴 a face that is bundled AND installed system-wide ({}) is NOT the bundle-only case — the "
          "exact situation that kept this suite red after Calibri was copied into ~/Library/Fonts "
          "(got {}: {})".format(_bf["both"], status, msg[:80]))


def _installed_face():
    """A face this machine resolves, preferring the one this skill recommends — or None."""
    ok = cfr._resolver()
    if ok is None:
        return None
    for preferred in ("Times New Roman", "Liberation Serif", "DejaVu Serif", "DejaVu Sans"):
        if ok(preferred) is True:
            return preferred
    return cfr.an_installed_face()


# 🔴 The property under test is "a face this machine really has passes clean" — NOT "Times New
# Roman passes clean". Hardcoding the recommended face made this arm depend on which fonts the
# runner happens to carry: it went green on macOS and red on the Linux CI box, which has no Times
# New Roman. The resolver asked is the one item 10 itself uses.
_SYSTEM_FACE = _installed_face()
if _SYSTEM_FACE:
    status, msg = pf.item10_fonts(_deck([("A headline in a face this host has", {"latin": _SYSTEM_FACE}, None)]))
    check(status == "PASS",
          "a system-wide face ({}) passes clean — the check must stay quiet on the answer it "
          "recommends, or authors learn to ignore it (got {}: {})"
          .format(_SYSTEM_FACE, status, msg))
else:
    check(True, "no enumerable system face on this machine — the passes-clean arm is skipped, not "
                "faked (it runs wherever any font is installed, which is every real runner)")


# ------------------------------------------ the face that draws the CJK is the face that is checked
GOOD = "Good Latin"
UNRES = {"Zz CJK Face", "Zz Theme EA", "Zz Hans Face", "Zz Jpan Face", "Zz Major EA",
         "Zz Minor EA", "Zz Arabic Face"}
CJK12 = "给准备申请的人三扇门长什么"[:12]


def _both(prs):
    """(item10 status, item10 message, check_fonts_resolve findings, carried) on ONE deck."""
    buf = io.BytesIO()
    prs.save(buf)
    path = os.path.join(tempfile.mkdtemp(prefix="fontcjk-"), "d.pptx")
    with open(path, "wb") as fh:
        fh.write(buf.getvalue())
    with _Resolver(UNRES):
        status, msg = pf.item10_fonts(prs)
        finds, facts = cfr.check(path)
    return status, msg, finds, facts["by_slot"]


def _agree(label, prs, expect):
    """🔴 The two instruments must name the SAME unresolved faces — the defect was their disagreeing."""
    status, msg, finds, carried = _both(prs)
    flagged = {f for _s, f, _m in finds if f in UNRES}
    named_by_10 = {f for f in UNRES if f in msg}
    check(flagged == set(expect) and named_by_10 == set(expect),
          "{}: both instruments name exactly {} (check_fonts_resolve {}, item 10 {}: {})"
          .format(label, sorted(expect) or "nothing", sorted(flagged), status, sorted(named_by_10)))
    return status, msg, finds, carried


st, msg, finds, carried = _agree("a CJK run whose <a:ea> face resolves nowhere",
                                 _deck([(CJK12, {"latin": GOOD, "ea": "Zz CJK Face"}, "zh-CN")]),
                                 {"Zz CJK Face"})
check(carried.get("Zz CJK Face", {}).get("ea") == 12 and not carried.get(GOOD),
      "...and the characters are credited to the face that DRAWS them: 12 East-Asian to the ea face, "
      "none to the Latin face (got {})".format(carried))
check(st == "ADVISORY" and "CJK" in msg and "60%" in msg,
      "...and item 10 says what a stand-in does to CJK widths, with the measured figure")
check(any(s_ == "block" and f == "Zz CJK Face" for s_, f, _m in finds),
      "...and check_fonts_resolve treats it as a real body face (block severity)")
check(all("%%" not in m and "60% of the true width" in m for _s, f, m in finds if f == "Zz CJK Face"),
      "...and its message prints the figure as written — a clause passed as a percent-format "
      "ARGUMENT keeps a doubled percent sign, which is what shipped to both gate paths the first "
      "time ({!r})".format([m[-120:] for _s, f, m in finds if f == "Zz CJK Face"]))

_agree("🔴 an ENGLISH run whose <a:ea> slot names an unresolved face (EAFONT set, no Chinese text)",
       _deck([("Quarterly revenue grew in every region", {"latin": GOOD, "ea": "Zz CJK Face"}, None)]),
       set())

st, msg, finds, carried = _agree("a MIXED run", _deck([("Revenue 营收 up 51%", {"latin": GOOD, "ea": "Zz CJK Face"}, None)]),
                                 {"Zz CJK Face"})
check(carried["Zz CJK Face"]["ea"] == 2 and carried[GOOD]["latin"] == len("Revenue 营收 up 51%") - 2,
      "...splits by character: 2 to the ea face, the rest to the Latin face (got {})".format(carried))
check(any(s_ == "note" and f == "Zz CJK Face" for s_, f, _m in finds),
      "...and two CJK characters stay under the block floor, reported only")

_agree("CJK text with NO ea face on the run, inheriting the theme's minor ea face",
       _deck([(CJK12, {"latin": GOOD}, None)], theme={("minor", "ea"): "Zz Theme EA"}),
       {"Zz Theme EA"})
_agree("the same unresolved theme ea face on a deck with no CJK text draws nothing",
       _deck([("Only English words on this slide", {"latin": GOOD}, None)],
             theme={("minor", "ea"): "Zz Theme EA"}), set())
_agree("an Office-style theme: <a:ea> empty, the Chinese face named per SCRIPT and picked by lang",
       _deck([(CJK12, {"latin": GOOD}, "zh-CN")],
             theme={("minor", "script:Hans"): "Zz Hans Face", ("minor", "script:Jpan"): "Zz Jpan Face"}),
       {"Zz Hans Face"})
_agree("...and a Japanese run in the same theme picks the Jpan face, not the Hans one",
       _deck([("ひらがなとカタカナの文章です", {"latin": GOOD}, "ja-JP")],
             theme={("minor", "script:Hans"): "Zz Hans Face", ("minor", "script:Jpan"): "Zz Jpan Face"}),
       {"Zz Jpan Face"})
_agree("a `+mn-ea` pointer on the run resolves to the theme's minor ea face",
       _deck([(CJK12, {"latin": GOOD, "ea": "+mn-ea"}, None)], theme={("minor", "ea"): "Zz Theme EA"}),
       {"Zz Theme EA"})

_title = Presentation()
_ts = _title.slides.add_slide(_title.slide_layouts[0])
_ts.shapes.title.text_frame.text = CJK12
_patch_theme(_title, {("major", "ea"): "Zz Major EA", ("minor", "ea"): "Zz Minor EA"})
_agree("CJK in a TITLE placeholder inherits the MAJOR (heading) ea face, not the body one",
       _reopen(_title), {"Zz Major EA"})

st, msg, finds, carried = _agree("Korean hangul is East-Asian too",
                                 _deck([("안녕하세요반갑습니다", {"latin": GOOD, "ea": "Zz CJK Face"}, None)]),
                                 {"Zz CJK Face"})
check(carried["Zz CJK Face"]["ea"] == 10, "...all 10 syllables credited to the ea face ({})".format(carried))
st, msg, finds, carried = _agree("Arabic is drawn from <a:cs>",
                                 _deck([("مرحبا بالعالم الجميل", {"latin": GOOD, "cs": "Zz Arabic Face"}, None)]),
                                 {"Zz Arabic Face"})
check(carried["Zz Arabic Face"]["cs"] == len("مرحبا بالعالم الجميل".replace(" ", "")),
      "...Arabic letters to the cs face, the spaces to Latin ({})".format(carried))
_agree("...and an English run with an unresolved cs face draws nothing in it",
       _deck([("Plain English on this slide", {"latin": GOOD, "cs": "Zz Arabic Face"}, None)]), set())

# text the old python-pptx shape walk never reached: a table cell, a group, a field
_t = Presentation()
_ts = _t.slides.add_slide(_t.slide_layouts[6])
_tbl = _ts.shapes.add_table(1, 1, Inches(0.5), Inches(0.5), Inches(6), Inches(1)).table
_run(_tbl.cell(0, 0).text_frame.paragraphs[0], CJK12, {"latin": GOOD, "ea": "Zz CJK Face"})
_agree("🔴 CJK in a TABLE CELL — a GraphicFrame has no text_frame, so the old walk skipped it",
       _reopen(_t), {"Zz CJK Face"})

_g = Presentation()
_gs = _g.slides.add_slide(_g.slide_layouts[6])
_grp = _gs.shapes.add_group_shape()
_gtb = _grp.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6), Inches(1))
_run(_gtb.text_frame.paragraphs[0], CJK12, {"latin": GOOD, "ea": "Zz CJK Face"})
_agree("CJK inside a GROUP", _reopen(_g), {"Zz CJK Face"})

_f = Presentation()
_fs = _f.slides.add_slide(_f.slide_layouts[6])
_ftb = _fs.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6), Inches(1))
_fp = _ftb.text_frame.paragraphs[0]._p
_fld = etree.SubElement(_fp, qn("a:fld"), {"id": "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", "type": "slidenum"})
_frpr = etree.SubElement(_fld, qn("a:rPr"), {"lang": "zh-CN"})
etree.SubElement(_frpr, qn("a:ea"), {"typeface": "Zz CJK Face"})
etree.SubElement(_fld, qn("a:t")).text = "第一页页码字段的文字"
_agree("CJK in a FIELD (<a:fld>) — a template's Chinese footer or date renders from it",
       _reopen(_f), {"Zz CJK Face"})

# the real builder's path: deckkit stamps <a:ea> itself when EAFONT is set
_saved = (dk.EAFONT,)
try:
    dk.EAFONT = "Zz CJK Face"
    _bp = dk.blank_deck()
    _bs = dk.add_slide(_bp)
    dk.text(_bs, 0.6, 0.6, 8, 0.8, [[(CJK12, 20, dk.DEEP, True, False, dk.FONT)]])
    _agree("a deck BUILT by deckkit with EAFONT set to an unresolved face", _reopen(_bp), {"Zz CJK Face"})
    _bp2 = dk.blank_deck()
    _bs2 = dk.add_slide(_bp2)
    dk.text(_bs2, 0.6, 0.6, 8, 0.8, [[("An English deck with EAFONT set", 20, dk.DEEP, True, False, dk.FONT)]])
    _agree("...and an ENGLISH deck built with the same EAFONT is not blamed for it", _reopen(_bp2), set())
finally:
    dk.EAFONT = _saved[0]


# ---------------------------------------------------- and the rule is written where it is read
_gui = (ROOT / "references" / "font-guidance.md").read_text(encoding="utf-8")
check("APP BUNDLE" in _gui and "DFonts" in _gui,
      "`font-guidance.md` carries the app-bundle trap with the real path, so the reason survives "
      "the check")
check("ACADEMIC / LAB register" in _gui and "Times New Roman" in _gui,
      "...and the academic/lab/conference type default, which is the question that was never "
      "answered anywhere: a lab room expects a conference face, not a designer sans")
_pur = (ROOT / "references" / "design-by-purpose.md").read_text(encoding="utf-8")
check(_pur.count("- **Type:**") >= 4,
      "the four academic purposes carry a `Type:` line — they had Palette / Density / Layout / "
      "Icons / Signature and no type row at all, which is why the face was inherited from whatever "
      "the template profile happened to say")
check("Type by register" in _pur,
      "...pointing at one shared note, so the rule is stated once rather than drifting five ways")
_setup = (ROOT / "references" / "deck-setup.md").read_text(encoding="utf-8")
check("APP BUNDLE" in _setup and "preflight_check.py` item 10" in _setup,
      "`deck-setup.md` §Fonts — the file SKILL.md routes to before the first `set_palette` — "
      "carries both rules and names the check that enforces one of them")


# ------------------------------------------------- the check runs in a FRESH PROCESS, as it does
tmp = Path(tempfile.mkdtemp(prefix="fontverif-"))
dk.set_palette(font=_SYSTEM_FACE or "Times New Roman",
                mono=_SYSTEM_FACE or "Courier New")
prs = dk.blank_deck()
s = dk.add_slide(prs)
dk.text(s, 1, 1, 8, 1, [[("a headline", 24, dk.DEEP, True, False, dk.FONT)]])
dk.speaker_notes(s, "n")
deck = tmp / "d.pptx"
prs.save(str(deck))
out = subprocess.run([sys.executable, str(SCRIPTS / "preflight_check.py"), str(deck)],
                     capture_output=True, text=True)
if _SYSTEM_FACE:
    check("10. Hand-off: fonts" in out.stdout and "resolve locally" in out.stdout,
          "🔴 and it works from a FRESH PROCESS — the way preflight actually runs — on a deck set "
          "in a face this machine has ({})".format(_SYSTEM_FACE))
else:
    check("10. Hand-off: fonts" in out.stdout,
          "🔴 preflight still reaches item 10 from a FRESH PROCESS even with no enumerable face")

for line in ok:
    print("  ok   " + line)
for line in bad:
    print("  FAIL " + line)
print("\n{} passed, {} failed".format(len(ok), len(bad)))
sys.exit(1 if bad else 0)
