#!/usr/bin/env python3
"""The three template-branch holes, each pinned in BOTH directions.

Every one of these was found on a real institutional-template deck (LKEB/LUMC) that the pipeline
called clean, and every one of them is the same shape: a check that could not SEE the thing it was
supposed to judge, and therefore either asserted something false or went quiet.

  1. backing fill   an inherited <p:bg> that is a PICTURE, and brand furniture painted as a SHAPE
                    on the layout/master, were both invisible to `_backing_fill`. Text over them
                    resolved to the page colour — reported as white-on-white on a blue band (6
                    false findings per run), and, the other way round, DARK TEXT ON A DARK BAND
                    PASSED SILENTLY. The second direction is what this really buys, so it is
                    tested first and hardest.
  2. theme colours  template furniture is filled with `<a:schemeClr>`, and the master's `<p:clrMap>`
                    is frequently NOT the identity (measured: bg1->dk2, tx1->lt1 on LKEB). Reading
                    the scheme without the map returns the wrong end of the palette.
  3. hidden fills   a PowerPoint placeholder carries `<a:noFill/>` plus an `<a14:hiddenFill>` with a
                    real colour. A descendant search finds the hidden one and reports a confident
                    wrong backing.

Run: python3 tests/test_template_branch_gates.py
"""
from __future__ import annotations

import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "scripts"))

from lxml import etree                                                    # noqa: E402
from pptx import Presentation                                             # noqa: E402
from pptx.util import Inches, Pt                                          # noqa: E402

import lint_deck as L                                                     # noqa: E402

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

OK, BAD = [], []


def ck(cond, msg):
    (OK if cond else BAD).append(msg)
    print(("  ok   " if cond else "  ✗    ") + msg)


def _band_sp(x, y, w, h, *, srgb=None, scheme=None, hidden=None):
    """A layout/master rect: solid srgb, solid schemeClr, or noFill + an a14:hiddenFill."""
    if srgb is not None:
        fill = '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>' % srgb
    elif scheme is not None:
        fill = '<a:solidFill><a:schemeClr val="%s"/></a:solidFill>' % scheme
    else:
        fill = ('<a:noFill/><a:extLst><a:ext uri="{909E8E84-426E-40DD-AFC4-6F175D3DCCD1}">'
                '<a14:hiddenFill xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main">'
                '<a:solidFill><a:srgbClr val="%s"/></a:solidFill></a14:hiddenFill>'
                '</a:ext></a:extLst>' % hidden)
    xml = (
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:nvSpPr><p:cNvPr id="900" name="band"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>%s</p:spPr></p:sp>'
        % (int(x * 914400), int(y * 914400), int(w * 914400), int(h * 914400), fill))
    return etree.fromstring(xml)


def _deck(*, band=None, bg_blip=False):
    """A one-slide deck with a title, plus optional layout furniture / picture background."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    slide = prs.slides.add_slide(prs.slide_layouts[6])            # blank
    if band is not None:
        slide.slide_layout._element.find(P + "cSld").find(P + "spTree").append(band)
    if bg_blip:
        master = slide.slide_layout.slide_master
        cSld = master._element.find(P + "cSld")
        for old in cSld.findall(P + "bg"):
            cSld.remove(old)
        bg = etree.fromstring(
            '<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<p:bgPr><a:blipFill><a:blip r:embed="rIdX"/><a:stretch/></a:blipFill>'
            '<a:effectLst/></p:bgPr></p:bg>')
        cSld.insert(0, bg)
    return prs, slide


def _title(slide, rgb, *, x=0.6, y=0.15, w=7.0, h=0.8):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = "A title long enough to measure"
    r.font.size = Pt(28)
    r.font.bold = True
    from pptx.dml.color import RGBColor
    r.font.color.rgb = RGBColor.from_string(rgb)
    return tb


def _findings(prs):
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "t.pptx")
        prs.save(path)
        import contextlib
        import io
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            L.lint(path)
        return buf.getvalue()


def _backing_of(prs, slide):
    sw, sh = prs.slide_width / 914400, prs.slide_height / 914400
    bx = L._boxes(slide, sw, sh, slide_no=1)
    ch = L._layout_chrome(slide, sw, sh)
    for ti, b in enumerate(bx):
        if b["text"]:
            return L._backing_fill(bx, ti, chrome=ch)
    return "NO-TEXT"


print("— 1. the direction that was SILENTLY PASSING: dark text on a template's dark band")
prs, slide = _deck(band=_band_sp(0.5, 0.1, 8.0, 0.9, srgb="003C66"))
_title(slide, "1A2B3C")
got = _backing_of(prs, slide)
ck(got == "003C66",
   "dark text over a layout's dark band now resolves that band (%r) — before this it resolved the "
   "page and the contrast check passed a 1.2:1 pair" % got)

print("\n— 2. the direction that was reporting FALSELY: white text on a template's brand band")
prs, slide = _deck(band=_band_sp(0.5, 0.1, 8.0, 0.9, srgb="007CC2"))
_title(slide, "FFFFFF")
got = _backing_of(prs, slide)
ck(got == "007CC2", "white text over a blue layout band resolves the band, not white (%r)" % got)
txt = _findings(prs)
ck("INVISIBLE TEXT" not in txt,
   "…and no INVISIBLE TEXT finding is emitted for it (the 6-per-run false positive)")

print("\n— 3. a THEME-coloured band, through a NON-identity clrMap")
prs, slide = _deck(band=_band_sp(0.5, 0.1, 8.0, 0.9, scheme="accent1"))
_title(slide, "FFFFFF")
got = _backing_of(prs, slide)
ck(isinstance(got, str) and got not in (None, "UNKNOWN") and len(got) == 6,
   "a schemeClr band resolves through the master's clrMap to a real hex (%r)" % got)
master = slide.slide_layout.slide_master
res = L._theme_resolver(master)
ck(res("accent1") == got, "…and it is the value the theme's own clrScheme declares for that slot")
ck(res("phClr") is None, "a `phClr` placeholder colour resolves to nothing rather than a guess")

print("\n— 4. a TRANSFORMED theme colour must stay UNKNOWN, never the base slot")
sp = _band_sp(0.5, 0.1, 8.0, 0.9, scheme="accent1")
_sch = sp.find(P + "spPr").find(A + "solidFill").find(A + "schemeClr")
etree.SubElement(_sch, A + "lumMod").set("val", "75000")
prs, slide = _deck(band=sp)
_title(slide, "FFFFFF")
ck(_backing_of(prs, slide) == "UNKNOWN",
   "a lumMod-tinted schemeClr reads UNKNOWN — reporting the base slot would be confidently wrong")

print("\n— 5. an a14:hiddenFill must NOT be read as the shape's colour")
prs, slide = _deck(band=_band_sp(0.5, 0.1, 8.0, 0.9, hidden="FF0000"))
_title(slide, "FFFFFF")
got = _backing_of(prs, slide)
ck(got != "FF0000",
   "a placeholder with <a:noFill/> plus a hidden red fill does not report red (%r)" % got)

print("\n— 6. an inherited PICTURE background stops being assumed white")
prs, slide = _deck(bg_blip=True)
_title(slide, "FFFFFF")
txt = _findings(prs)
ck("INVISIBLE TEXT" not in txt,
   "white text on a master's blipFill background is no longer reported as white-on-white")
sw, sh = prs.slide_width / 914400, prs.slide_height / 914400
rec = L._slide_bg_box(slide, sw, sh)
ck(rec is not None and rec.get("unk") is True,
   "…because the inherited picture background is recorded as UNKNOWN, which is what routes the "
   "question to the render-pixel check instead of silencing it")

print("\n— 7. the case the old comment protects must STILL return None")
prs, slide = _deck()                                              # python-pptx default master
rec = L._slide_bg_box(slide, sw, sh)
ck(rec is None,
   "a bare theme <p:bgRef> still returns None — it is a reference, not evidence that anything is "
   "painted, and an `unk` record there would poison every plain deck")

print("\n— 8. chrome must not leak into the content checks")
prs, slide = _deck(band=_band_sp(0.5, 0.1, 8.0, 0.9, srgb="007CC2"))
_title(slide, "FFFFFF")
bxs = L._boxes(slide, sw, sh, slide_no=1)
ck(all(b.get("st") != "AUTO_SHAPE" or b.get("bg") or b["w"] < 7.9 for b in bxs),
   "the layout band is NOT injected into `bx` — coverage, counts and overlap checks all filter "
   "only `not s['bg']`, so six inherited rects per slide would move all of them at once")

print()
print("%d passed, %d failed" % (len(OK), len(BAD)))
raise SystemExit(1 if BAD else 0)
