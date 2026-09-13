#!/usr/bin/env python3
"""Two escape hatches an author could not find, and a default that collided with another default.

🔴 `adopt()` — python-pptx's `add_shape` leaves the theme `<p:style>` on a shape, so LibreOffice
draws a soft drop shadow under it and `lint_layout` reports `INHERITED_EFFECT`. The only cure was
`_flat`, a PRIVATE function. Measured on a real build: the author hit the finding, went reading
the source, and called the underscore name from a deck script. Every primitive deckkit lacks
arrives through `add_shape`, so this is part of the contract, not an implementation detail.

🔴 `motif_legend`'s default placement collided with `bottom_callout`'s. Both anchor to the bottom
of `content_band`, so two helpers left on their defaults drew into the same strip and the build
died on TEXT_OVERLAP with nothing naming the cause.

Not a defect, recorded so it is not "fixed" later: `motif_legend` has ALWAYS taken `x=`/`y=`. The
build that collided passed neither, having skipped the one `sigs.py` lookup that would have shown
them. The default is what changed; the parameters were there.
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import deckkit as dk  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches  # noqa: E402

fails: list[str] = []


def check(cond, msg):
    if not cond:
        fails.append(msg)


dk.set_palette(font="Helvetica Neue")
W, H = 13.333, 7.5

# ── adopt() flattens a hand-added shape, and is PUBLIC ──────────────────────────────────────────
_adopt = getattr(dk, "adopt", None)
check(callable(_adopt),
      "dk.adopt is gone — the only public route for a shape deckkit has no helper for")
if callable(_adopt):
    check(not _adopt.__name__.startswith("_"), "adopt is private again")

prs = dk.blank_deck(W, H)
s = dk.add_slide(prs)
raw = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(1), Inches(1), Inches(1), Inches(1))
check("<p:style>" in raw._element.xml,
      "a raw add_shape no longer carries <p:style>; if python-pptx changed, adopt's reason needs "
      "re-checking rather than assuming")
if callable(_adopt):
    out = _adopt(raw)
    check("<p:style>" not in raw._element.xml, "adopt did not strip the inherited theme style")
    check(out is raw, "adopt must return the shape, so it composes: "
                      "dk.adopt(shapes.add_shape(...))")

# ── the two defaults no longer collide ──────────────────────────────────────────────────────────
prs2 = dk.blank_deck(W, H)
s2 = dk.add_slide(prs2)
dk.bottom_callout(s2, 0.9, 11.5, "NOTE", "a bottom callout that anchors to the band")
dk.motif_legend(s2, "the ring means working")
try:
    dk.lint_layout(prs2, strict=True)
except RuntimeError as exc:                                           # noqa: BLE001
    fails.append("a legend and a callout on one slide, both on their defaults, still collide: {}"
                 .format(str(exc)[:120]))

# ── the legend still lands where it is TOLD to ──────────────────────────────────────────────────
prs3 = dk.blank_deck(W, H)
s3 = dk.add_slide(prs3)
tb = dk.motif_legend(s3, "the ring means working", x=2.0, y=3.0)
from pptx.util import Emu  # noqa: E402
got = None
for sh in s3.shapes:
    if getattr(sh, "has_text_frame", False) and "the ring means" in sh.text_frame.text:
        got = (Emu(sh.left).inches, Emu(sh.top).inches)
check(got is not None, "the legend text is not on the slide at all")
if got:
    check(abs(got[1] - 3.0) < 0.15,
          "an explicit y= was overridden by the collision-avoidance default: asked 3.00, got "
          "{:.2f} — the new default must not fight an author who said where to put it"
          .format(got[1]))

# ── and a legend ALONE still sits in the band, not off the page ─────────────────────────────────
prs4 = dk.blank_deck(W, H)
s4 = dk.add_slide(prs4)
dk.motif_legend(s4, "the ring means working")
ys = [Emu(sh.top).inches for sh in s4.shapes
      if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
check(ys and 0.0 < min(ys) < H, "a lone legend landed off the canvas: {}".format(ys))

print("\n".join("FAIL " + f for f in fails) if fails else "", end="")
print("[test_adopt_and_legend_default] {}".format(
    "FAILED: {} problem(s)".format(len(fails)) if fails else "ok"))
sys.exit(1 if fails else 0)
