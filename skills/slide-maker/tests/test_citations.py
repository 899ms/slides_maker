#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Every citation this skill had ever built was typed by hand, and nothing checked it against a .bib.

Measured before `scripts/citations.py`: `sources_page` renders a list of STRINGS and `source_note`
renders a provenance line of STRINGS. Nothing read a bibliography, nothing linked an in-text marker
to a reference list, and nothing could tell a correct year from a remembered one. Retyping is where
a year drifts by one and a middle author disappears — and on a slide, a misattributed result is a
claim about a real person's work.

What is pinned here:
  * the PARSER, against the shapes a real bibliography has (braced titles, LaTeX accents, quoted
    and parenthesised entries, particled surnames) rather than against a tidy fixture;
  * the REFUSAL — no author, no title or no year raises instead of printing `n.d.`, which is this
    skill's never-invent floor applied to the one place a fabrication is unfalsifiable;
  * the GATE, against a REAL BUILT DECK: a dangling marker, a cited key that is on no slide, an
    entry nothing cites, and both NOT CHECKED refusals.

Run: python3 tests/test_citations.py
"""
from __future__ import annotations

import pathlib
import re
import sys
import tempfile
import warnings

HERE = pathlib.Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL / "scripts"))
warnings.simplefilter("ignore")

import check_citations as cc                                              # noqa: E402
import citations as cit                                                   # noqa: E402
import deckkit as dk                                                      # noqa: E402
from pptx import Presentation                                             # noqa: E402
from pptx.dml.color import RGBColor                                       # noqa: E402

OK, BAD = [], []
TMP = pathlib.Path(tempfile.mkdtemp(prefix="citations-"))
INK = RGBColor(0x22, 0x22, 0x22)

BIB = r"""
% a comment line, and a @string nobody should read as an entry
@string{mrm = "Magnetic Resonance in Medicine"}
@article{lustig2007,
  author  = {Lustig, Michael and Donoho, David and Pauly, John M.},
  title   = {Sparse {MRI}: The application of compressed sensing for rapid {MR} imaging},
  journal = {Magnetic Resonance in Medicine}, year = {2007}, doi = {10.1002/mrm.21391}}
@inproceedings{muller2019,
  author    = {M{\"u}ller, Hans and van der Berg, Anna and Krej{\v{c}}{\'\i}, Jan},
  title     = {Sampling and reconstruction},
  booktitle = "Information Processing in Medical Imaging", year = 2019}
@misc(zbontar2018,
  author = {Zbontar, Jure and Knoll, Florian},
  title  = {fastMRI: An open dataset},
  eprint = {1811.08839}, year = {2018}, url = {https://arxiv.org/abs/1811.08839})
@article{noyear, author = {Somebody, A.}, title = {A paper with no year}, journal = {J.}}
@book{solo, author = {Tufte, Edward R.}, title = {The Visual Display of Quantitative Information},
  publisher = {Graphics Press}, year = {2001}}
"""


def ck(cond, msg, detail=""):
    (OK if cond else BAD).append(msg)
    print(("  ok   " if cond else "  ✗    ") + msg + (("  — %s" % (detail,)) if detail and not cond else ""))


print("— the parser, against the shapes a real bibliography has")
DB = cit.parse_bibtex(BIB)
ck(sorted(DB) == ["lustig2007", "muller2019", "noyear", "solo", "zbontar2018"],
   "five entries parse and @string/@comment are not among them — a `@string` read as an entry "
   "becomes a phantom reference nobody can find", sorted(DB))
ck(DB["lustig2007"]["title"] == "Sparse MRI: The application of compressed sensing for rapid MR imaging",
   "🔴 a title with PROTECTIVE braces survives whole — a regex over `field = {value}` stops at the "
   "first inner `}`, i.e. at `{MRI}`, and truncates a title on a real .bib while passing a tidy "
   "fixture", DB["lustig2007"]["title"])
ck(DB["muller2019"]["booktitle"] == "Information Processing in Medical Imaging",
   "a QUOTED value parses like a braced one", DB["muller2019"].get("booktitle"))
ck(DB["zbontar2018"]["year"] == "2018" and DB["muller2019"]["year"] == "2019",
   "a PARENTHESISED entry and a BARE numeric year both parse")
ck("Müller" in DB["muller2019"]["author"] and "Krejčí" in DB["muller2019"]["author"],
   "LaTeX accents resolve, including one written `{\\v{c}}` — a name is not decoration",
   DB["muller2019"]["author"])
ck(cit._deaccent(r"Nov\'{a}k-\r{a}\q{x}") == "Novák-åx",
   "...and an accent the table does NOT know keeps its base letter rather than printing `\\q{x}`",
   cit._deaccent(r"Nov\'{a}k-\r{a}\q{x}"))
ck(len(cit.authors(DB["muller2019"])) == 3 and len(cit.authors(DB["lustig2007"])) == 3,
   "authors split on ` and `, never on the comma inside `Last, First`",
   cit.authors(DB["muller2019"]))
ck(cit.authors({"author": "Sampling and reconstruction group"})[0] == "Sampling",
   "(the same rule means a literal ' and ' inside a name field splits too — BibTeX's own rule, "
   "which is why a corporate author is braced in the .bib)")

print("\n— names: the part everyone gets wrong")
for written, want in (("van der Berg, Anna", "van der Berg"), ("Anna van der Berg", "van der Berg"),
                      ("Lustig, Michael", "Lustig"), ("John M. Pauly", "Pauly"),
                      ("Tufte, Edward R.", "Tufte"), ("张三", "张三")):
    ck(cit.surname(written) == want, "surname(%r) -> %r" % (written, want), cit.surname(written))
ck(cit.initials("Pauly, John M.") == "J. M." and cit.initials("John M. Pauly") == "J. M.",
   "initials come out the same from both spellings", cit.initials("John M. Pauly"))
ck(cit.initials("Lustig") == "",
   "a surname-only name yields NO initial — an invented first initial is an invented fact")

print("\n— 🔴 the refusal: it does not fill in what the bibliography does not have")
ck(cit.missing(DB["noyear"]) == ["year"], "the one missing field is named", cit.missing(DB["noyear"]))
for key, field in (("noyear", "year"),):
    try:
        cit.format_reference(DB[key], "numeric", 1)
        ck(False, "%r is REFUSED, not printed with a placeholder" % key, "it formatted anyway")
    except cit.Incomplete as exc:
        ck(key in str(exc) and field in str(exc) and "never-invent" in str(exc).replace("-", "-"),
           "%r is REFUSED and the message names the key and the field, so the author can go get it"
           % key)
for bad_entry in ({"key": "x", "title": "T", "year": "2020"},          # no author
                  {"key": "x", "author": "A, B", "year": "2020"},      # no title
                  {"key": "x", "author": "A, B", "title": "T"}):       # no year
    try:
        cit.format_reference(bad_entry, "author-year")
        ck(False, "an entry missing one of author/title/year is refused", bad_entry)
    except cit.Incomplete:
        ck(True, "an entry missing %s is refused"
           % ("author" if "author" not in bad_entry else "title" if "title" not in bad_entry else "year"))

print("\n— the two styles, derived from the entry")
ck(cit.in_text(DB["lustig2007"], "numeric", 3) == "[3]", "numeric marker is its position")
try:
    cit.in_text(DB["lustig2007"], "numeric")
    ck(False, "numeric style without a position is refused", "it invented one")
except ValueError:
    ck(True, "numeric style without a position is REFUSED — a marker guessed from nothing points "
             "at the wrong paper, which is worse than no marker")
ck(cit.in_text(DB["solo"], "author-year") == "(Tufte, 2001)", "one author",
   cit.in_text(DB["solo"], "author-year"))
ck(cit.in_text(DB["zbontar2018"], "author-year") == "(Zbontar & Knoll, 2018)", "two authors",
   cit.in_text(DB["zbontar2018"], "author-year"))
ck(cit.in_text(DB["muller2019"], "author-year") == "(Müller et al., 2019)", "three authors",
   cit.in_text(DB["muller2019"], "author-year"))
ck(cit.in_text(DB["muller2019"], "author-year", etal="等", amp="与") == "(Müller等, 2019)",
   "🔴 a Chinese deck gets 等/与 rather than English inside a Chinese sentence",
   cit.in_text(DB["muller2019"], "author-year", etal="等", amp="与"))
num = cit.format_reference(DB["lustig2007"], "numeric", 1)
ck(num.startswith("[1] M. Lustig, D. Donoho, J. M. Pauly,") and '"Sparse MRI' in num
   and "Magnetic Resonance in Medicine, 2007." in num and "doi:10.1002/mrm.21391" in num,
   "the numeric line carries authors, title, venue, year and DOI, in that order", num)
ay = cit.format_reference(DB["lustig2007"], "author-year", 1)
ck(ay.startswith("Lustig, M., Donoho, D. & Pauly, J. M. (2007).") and ay.endswith("mrm.21391"),
   "the author-year line reads as a reference list entry", ay)
ck("arXiv:1811.08839" in cit.format_reference(DB["zbontar2018"], "numeric", 1),
   "an arXiv entry with no journal falls back to its eprint, rather than to a blank venue")
ck("Graphics Press" in cit.format_reference(DB["solo"], "numeric", 1),
   "a book falls back to its publisher — the venue chain is a chain, not one field")
ck(cit.doi_url(DB["lustig2007"]) == "https://doi.org/10.1002/mrm.21391"
   and cit.doi_url(DB["zbontar2018"]) == "https://arxiv.org/abs/1811.08839"
   and cit.doi_url(DB["solo"]) is None,
   "a DOI becomes a URL, a bare url field is used as-is, and an entry with neither links nowhere")

print("\n— the marker scanner, on text that is not a fixture")
ck(cc.numeric_markers("as in [3], [5, 7] and [10-12]") == {3, 5, 7, 10, 11, 12},
   "single, list and RANGE markers", cc.numeric_markers("as in [3], [5, 7] and [10-12]"))
ck(cc.numeric_markers("we enrolled 14 patients in 2020") == set(),
   "plain numbers are NOT read as citations")
ck(cc.numeric_markers("95% CI [0.2, 0.4]") == set(),
   "a decimal confidence interval is NOT read as a citation — the common form on a clinical slide")
ck(cc.numeric_markers("range [12, 34] mm") == {12, 34},
   "...but a bracketed INTEGER range IS, and that ambiguity is real: in numeric style `[12, 34]` "
   "is how two papers are cited. The gate says so in the finding and the waiver is the escape — "
   "silently ignoring plausible markers would delete the check this file exists for")
got = cc.author_year_markers("as (Lustig et al., 2007) showed, and Zbontar (2018) confirmed")
ck(got == {("lustig", "2007"), ("zbontar", "2018")},
   "🔴 both marker shapes are read — parenthetical AND narrative. The scanner lowercases for "
   "comparison, and lowercasing BEFORE the match deleted every narrative marker, because that one "
   "is recognised by its capital", got)
ck(cc.author_year_markers("（张三等，2020）") == {("张三", "2020")},
   "a full-width CJK marker is read — 等 takes no space before it, so an `\\s+` split leaves 张三等",
   cc.author_year_markers("（张三等，2020）"))

print("\n— the gate, against a REAL built deck")
bib_path = TMP / "refs.bib"
bib_path.write_text(BIB, encoding="utf-8")
KEYS = ["lustig2007", "muller2019", "zbontar2018"]
ENTRIES = [DB[k] for k in KEYS]


def build(name, body_text, *, entries=ENTRIES, style="numeric"):
    prs = dk.blank_deck()
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    dk.text(s1, 0.7, 0.7, 8.6, 1.0, [[("Where it stands", 26, INK, True, False, dk.FONT)]])
    dk.text(s1, 0.7, 2.0, 8.6, 2.0, [[(body_text, 15, INK, False, False, dk.FONT)]])
    if entries:
        cit.reference_page(prs.slides.add_slide(prs.slide_layouts[6]), entries, style=style)
    p = TMP / name
    prs.save(str(p))
    return str(p)


PLAN = {"bib": "refs.bib", "style": "numeric", "keys": KEYS}
clean = build("clean.pptx", "CS set the baseline [1]; sampling matters [2]; the benchmark is open [3].")
finds, facts = cc.check(clean, PLAN, root=str(TMP))
ck(finds == [] and facts["list_slides"] == [2],
   "every marker resolves, every entry is cited, and the reference list is FOUND on slide 2", finds)

finds, _f = cc.check(build("dangle.pptx", "CS set the baseline [1]; and see [7]."), PLAN, root=str(TMP))
ck(any(s == "block" and c == "DANGLING MARKER" and "[7]" in m for s, c, m in finds),
   "🔴 [7] over a three-entry list BLOCKS — the leftover of a cut slide, and the one an audience "
   "member actually looks up", finds)
ck(any(c == "UNCITED" and s == "note" for s, c, m in finds),
   "...and the entries nobody cited are reported as notes in the same pass, not one per run", finds)

finds, _f = cc.check(clean, dict(PLAN, keys=KEYS + ["solo"]), root=str(TMP))
ck(any(s == "block" and c == "NOT IN THE LIST" and "solo" in m for s, c, m in finds),
   "a cited key that is on NO slide blocks — the marker points at a page that is not in the deck",
   finds)
finds, _f = cc.check(clean, dict(PLAN, keys=KEYS + ["ghost2099"]), root=str(TMP))
ck(any(s == "block" and c == "NO SUCH ENTRY" for s, c, m in finds),
   "a key the bibliography does not have blocks — nobody in the room can look it up either", finds)
finds, _f = cc.check(clean, dict(PLAN, keys=["noyear"]), root=str(TMP))
ck(any(s == "block" and c == "INCOMPLETE ENTRY" and "year" in m for s, c, m in finds),
   "an entry with no year blocks, naming the field — the line cannot be built without inventing it",
   finds)

ay_deck = build("ay.pptx", "As (Lustig et al., 2007) and (Müller et al., 2019) showed.",
                entries=[DB["lustig2007"], DB["muller2019"]], style="author-year")
finds, _f = cc.check(ay_deck, {"bib": "refs.bib", "style": "author-year",
                               "keys": ["lustig2007", "muller2019"]}, root=str(TMP))
ck(finds == [], "the author-year style checks end to end on its own built deck", finds)
finds, _f = cc.check(build("ay2.pptx", "As (Lustig et al., 2007) and (Nobody et al., 1999) showed.",
                           entries=[DB["lustig2007"]], style="author-year"),
                     {"bib": "refs.bib", "style": "author-year", "keys": ["lustig2007"]},
                     root=str(TMP))
ck(any(s == "block" and c == "DANGLING MARKER" and "nobody" in m.lower() for s, c, m in finds),
   "...and an author-year marker in no entry blocks too — the check is not numeric-only", finds)

print("\n— what it refuses to answer, the waiver, and the path rule")
for plan, why in ((None, "no citation plan recorded"),
                  ({"bib": "refs.bib", "keys": []}, "a plan with no keys")):
    try:
        cc.check(clean, plan and cc.recorded_citations({"content": {"citations": plan}}), root=str(TMP))
        ck(False, "NOT CHECKED: %s" % why, "it returned a verdict")
    except RuntimeError as exc:
        ck("content.citations" in str(exc),
           "NOT CHECKED: %s — and the message names the field that would let it answer" % why)
try:
    cc.check(clean, {"bib": "../../../etc/passwd", "keys": ["x"]}, root=str(TMP))
    ck(False, "a bibliography path escaping the deck folder is refused", "it read it")
except RuntimeError as exc:
    ck("outside the deck folder" in str(exc),
       "🔴 a .bib path escaping the deck folder is REFUSED — the record is written by a build, and "
       "a build's record is not a licence to read anywhere on the machine", exc)
_fw, _fa = cc.check(clean, dict(PLAN, keys=KEYS + ["ghost2099"]), root=str(TMP),
                    waive="ghost2099 is cited from the appendix deck, which ships separately")
ck(_fa.get("waived") and [f for f in _fw if f[0] == "block"],
   "a written waiver does not delete the finding — it is recorded beside it")

print("\n— the rendered page: one typography, whatever the .bib happened to carry")
prs = Presentation(clean)
runs = [r for sh in list(prs.slides)[1].shapes if sh.has_text_frame
        for p in sh.text_frame.paragraphs for r in p.runs]
linked = [r for r in runs if r.hyperlink.address]
ck(len(linked) >= 2, "the entries that have a DOI or URL are LINKED — a reference one click from "
                     "the paper", len(linked))
ck(all(r.font.underline is False for r in linked),
   "...and the link carries no underline: on a reference list the DOI text is the affordance")
theme = [p for p in prs.part.package.iter_parts() if "theme" in str(p.partname)]
_hl = re.search(r'<a:hlink><a:srgbClr val="([0-9A-Fa-f]{6})"',
                theme[0].blob.decode("utf-8", "replace")) if theme else None
_want = "%02X%02X%02X" % (dk.DEEP[0], dk.DEEP[1], dk.DEEP[2])
ck(_hl is not None and _hl.group(1).upper() == _want,
   "🔴 the DECK's hyperlink colour is its own INK (%s), not the Office default. A linked run "
   "carries an explicit solidFill and a renderer paints it in the THEME's hlink colour anyway — "
   "measured, one references page came out half navy and half Word-blue, decided by which entries "
   "happened to have a DOI. (Asserting merely that SOME srgbClr sits there passes on an untouched "
   "deck — the stock theme ships 0563C1 — and this check WAS that until a mutant proved it.)"
   % _want, _hl and _hl.group(1))

print("\n— both runtimes run it, and the record carries the plan across")
import check_gate_parity as gp                                            # noqa: E402
ck("citations" in gp.RECORD_FED,
   "citations is declared RECORD-FED, so parity demands tests/test_schema_reach.py prove both "
   "runtimes' records can be READ")
_shared = (SKILL / "scripts" / "render_deck.py").read_text(encoding="utf-8")
_codex = (SKILL / "scripts" / "codex_delivery_gate.py").read_text(encoding="utf-8")
ck("_gate_section('citations')" in _shared and "def check_citations" in _codex,
   "the gate is wired on the shared path and the Codex path")
for name, src in (("shared", _shared), ("codex", _codex)):
    ck("reference_page" in src and "content.citations" in src.replace('"content": {"citations"', "content.citations"),
       "%s: its message names the field and the helper that renders the list" % name)
for name, f in (("shared", "deck_gates.py"), ("codex", "codex_delivery_gate.py")):
    src = (SKILL / "scripts" / f).read_text(encoding="utf-8")
    ck('"citations"' in src and "check_citations" in src,
       "%s: the --init scaffold SHOWS the field and names the gate that reads it" % name)
import sigs                                                               # noqa: E402
ck("citations" in sigs.MODULES and "reference_page" in sigs.EXAMPLES,
   "`sigs.py reference_page --example` hands back a runnable call — the citation helpers are found "
   "through the same door as every other component, and the smoke suite EXECUTES that scaffold")

print()
print("%d passed, %d failed" % (len(OK), len(BAD)))
raise SystemExit(1 if BAD else 0)
