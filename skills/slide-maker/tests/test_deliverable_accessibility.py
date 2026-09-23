#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""The PDF that travels, the handout nobody had, and the tally that separates "passed" from "checked".

Three things this suite pins, each found by looking at the artifact rather than at the code.

🔴 THE DELIVERED PDF DROPPED EVERY ALT TEXT. The a11y gate spends its existence demanding image
descriptions, and a plain `--convert-to pdf` exports a deck that is *tagged* — `/StructTreeRoot`,
`/Marked true`, `/Lang` — and carries NOT ONE `/Alt` value. Every description written for a screen
reader stopped at the .pptx, in the one artifact that actually gets emailed. With LibreOffice's
`PDFUACompliance` filter the same deck exports with `pdfuaid` and the descriptions present, the
raster is byte-identical, and the file is ~100 bytes bigger. WCAG 2.1 AA became the ADA Title II
standard in April 2026, which is what makes this a floor for the academic and public-sector decks
this skill is used for rather than a nicety.

🔴 THE SPEAKER HANDOUT DID NOT EXIST. The skill requires speaker notes, reads them to estimate the
talk's length, and then handed over a .pptx and a slide PDF in which the notes are invisible — the
one artifact a presenter rehearses from was never produced.

🔴 "PASSED" WAS NOT "CHECKED". A gate that could not bind printed its own NOT-CHECKED line and
nothing counted them, so a run ended with "all hand-off gates pass" whether 24 sections bound or
15 did. Measured in one session's audits: three gates were silent on every deck because nobody
asked for the field they read, and the template gate could bind to 1 of 11 registered templates.
Each was found by a human reading a transcript. The ledger makes the tool report it.

Run: python3 tests/test_deliverable_accessibility.py
"""
from __future__ import annotations

import os
import pathlib
import re
import subprocess
import sys
import tempfile
import warnings

HERE = pathlib.Path(__file__).resolve().parent
SKILL = HERE.parent
sys.path.insert(0, str(SKILL / "scripts"))
warnings.simplefilter("ignore")

import deckkit as dk                                                      # noqa: E402
import render_deck as rd                                                  # noqa: E402
from pptx.dml.color import RGBColor                                       # noqa: E402

OK, BAD = [], []
TMP = pathlib.Path(tempfile.mkdtemp(prefix="deliverable-"))
ALT = "Line chart: SSIM rises from 0.81 to 0.93 as the cascade deepens"


def ck(cond, msg, detail=""):
    (OK if cond else BAD).append(msg)
    print(("  ok   " if cond else "  ✗    ") + msg + (("  — %s" % (detail,)) if detail and not cond else ""))


def _deck(name, *, notes=True, alt=True):
    prs = dk.blank_deck()
    for i in range(1, 4):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text_frame.text = "Slide %d" % i
        dk.text(s, 0.7, 2.0, 8.6, 1.0,
                [[("Body copy on slide %d." % i, 15, dk.DEEP, False, False, dk.FONT)]])
        if alt and i == 1:
            try:
                from PIL import Image
                png = TMP / "fig.png"
                if not png.exists():
                    Image.new("RGB", (400, 300), (0x1C, 0x5D, 0x99)).save(png)
                pic = dk.picture(s, str(png), 5.6, 2.6, 3.2, 2.4)
                pic._element._nvXxPr.cNvPr.set("descr", ALT)
            except ImportError:
                pass
        if notes:
            dk.speaker_notes(s, "Say the thing that slide %d does not say: the reason it matters, "
                                "and what the room has to decide about it." % i)
    path = TMP / name
    prs.save(str(path))
    return str(path)


SOFFICE = rd.find_soffice()
print("— the PDF that travels")
if not SOFFICE:
    ck(True, "(LibreOffice absent on this host — the export checks are skipped, not silently passed)")
else:
    deck = _deck("a11y.pptx")
    out = tempfile.mkdtemp(dir=str(TMP))
    # 🔴 The RENDER's export is the rasterization intermediate and must stay a plain one. The first
    # version put the PDF/UA filter here because the page pixmap hashed identically on this host —
    # and CI, on another LibreOffice, stopped raising `CAPTION NOT ALIGNED`, a PIXEL check reading
    # those PNGs, on a deck built to trip it. "Byte-identical" was true of one host and asserted of
    # all of them.
    _plain_pdf, _plain_res, _plain_cmd = rd._render_pdf(SOFFICE, deck, out)
    ck("PDFUACompliance" not in " ".join(_plain_cmd),
       "🔴 the RENDER exports plainly — the accessibility filter is off the path that feeds the "
       "PNGs, so no pixel gate can ever change because of it", _plain_cmd)
    dest_pdf = str(TMP / "a11y-deliverable.pdf")
    pdf = rd.render_accessible_pdf(SOFFICE, deck, dest_pdf)
    raw = pathlib.Path(pdf).read_bytes() if pdf else b""
    ck(bool(pdf) and raw, "the DELIVERABLE export produced a PDF", pdf)
    ck(b"/StructTreeRoot" in raw and b"/Marked true" in raw,
       "it is TAGGED — a structure tree, which a `Print to PDF` never has")
    ck(b"pdfuaid" in raw,
       "🔴 ...and it DECLARES PDF/UA, which the plain export does not. The declaration is what a "
       "checker and a procurement reviewer look for")
    hexed = "".join("%04X" % ord(c) for c in ALT[:20]).encode()
    ck(hexed in raw or ALT.encode() in raw or ALT.encode("utf-16-be") in raw,
       "🔴 ...and the picture's ALT TEXT is in it. Before the filter the same deck exported with a "
       "structure tree and zero /Alt values: every description the a11y gate demanded stopped at "
       "the .pptx", raw[:0])
    ck("PDFUACompliance" in rd.PDF_UA_FILTER and "UseTaggedPDF" in rd.PDF_UA_FILTER,
       "the filter names both options, so a LibreOffice that honours one and not the other still "
       "tags the file")
    ck("render_accessible_pdf(soffice, pptx, pdf_dest)" in
       (SKILL / "scripts" / "render_deck.py").read_text(encoding="utf-8"),
       "...and the hand-off run INVOKES it on the delivered copy — an export nothing calls is an "
       "export nobody receives")

print("\n— the speaker handout")
if not SOFFICE:
    ck(True, "(LibreOffice absent — skipped)")
else:
    noted = _deck("talk.pptx")
    dest = str(TMP / "talk-notes.pdf")
    got = rd.render_notes_pdf(SOFFICE, noted, dest)
    ck(got == dest and os.path.exists(dest),
       "a deck WITH notes gets a handout beside it — the artifact a presenter rehearses from, "
       "which this skill demanded the notes for and never produced", got)
    if got:
        try:
            import fitz
            doc = fitz.open(dest)
            ck(doc.page_count == 3, "one page per slide", doc.page_count)
            txt = re.sub(r"\s+", " ", doc[0].get_text())   # the renderer wraps mid-phrase
            ck("Slide 1" in txt and "what the room has to decide" in txt,
               "🔴 each page carries the slide AND its notes — the notes are the half the slide PDF "
               "cannot show, which is the whole reason this file exists")
            doc.close()
        except ImportError:
            ck(True, "(pymupdf absent — the page-level read is skipped)")
    # …and the --deliverables path must CALL it. That run is gated behind the quality record, so
    # this reads the call site: a renderer nothing invokes is a renderer nobody gets.
    ck("render_notes_pdf(soffice, pptx, notes_dest)" in
       (SKILL / "scripts" / "render_deck.py").read_text(encoding="utf-8"),
       "the hand-off run INVOKES the handout renderer — testing the function alone would pass "
       "with the call site deleted, which is how a feature ships unreachable")
    bare = _deck("nonotes.pptx", notes=False)
    dest2 = str(TMP / "nonotes-notes.pdf")
    ck(rd.render_notes_pdf(SOFFICE, bare, dest2) is None and not os.path.exists(dest2),
       "a deck with NO notes gets no handout rather than three empty pages — and says so")

print("\n— the coverage ledger: 'passed' and 'checked' are different sentences")
ck(hasattr(rd, "not_checked") and hasattr(rd, "coverage_ledger") and hasattr(rd, "print_coverage_ledger"),
   "the shared path has a ledger at all")
_src = (SKILL / "scripts" / "render_deck.py").read_text(encoding="utf-8")
# The ledger's own header prints the words "NOT CHECKED" — it is the tally, not a gate reporting
# itself, so it is named here rather than quietly widening the pattern until nothing matches.
_LEDGER_OWN = "COVERAGE: {} section(s) ran"
_bare = [l for l in _src.splitlines()
         if ("NOT CHECKED" in l or "not applied" in l) and re.search(r"\bprint\(", l)
         and _LEDGER_OWN not in l]
ck(_bare == [],
   "🔴 EVERY not-checked line goes through the ledger — one left on a bare `print` is a gate that "
   "silently drops out of the tally, which is the exact shape of the problem the ledger exists "
   "for", _bare[:2])
rd._NOT_CHECKED.clear()
rd._SECTIONS_RUN.clear()
for name in ("alpha", "beta", "gamma"):
    if name not in rd._SECTIONS_RUN:
        rd._SECTIONS_RUN.append(name)
rd.not_checked("  [--] beta: NOT CHECKED — nobody recorded the field", section="beta")
checked, rows = rd.coverage_ledger(rd._SECTIONS_RUN)
ck((checked, [r[0] for r in rows]) == (2, ["beta"]),
   "the tally counts what bound and names what did not", (checked, rows))
ck(rows and "nobody recorded the field" in rows[0][1] and "beta:" not in rows[0][1],
   "...and the row carries the REASON with its chrome and restated label stripped, so the "
   "actionable half is not pushed off the end of the line", rows)
_codex = (SKILL / "scripts" / "codex_delivery_gate.py").read_text(encoding="utf-8")
_bare_cx = [l for l in _codex.splitlines()
            if "NOT CHECKED" in l and re.search(r"\bprint\(", l) and "def " not in l]
ck(_bare_cx == [],
   "no not-checked line on the Codex path bypasses its ledger — a runtime that reports coverage "
   "and one that does not would disagree about what a clean run means", _bare_cx[:2])
import importlib.util                                                     # noqa: E402
_spec = importlib.util.spec_from_file_location("cdg_for_test",
                                               SKILL / "scripts" / "codex_delivery_gate.py")
_cdg = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_cdg)
_cdg._NOT_CHECKED.clear()
_cdg.not_checked("  [--] talk time NOT CHECKED — no budget was recorded (not clean)")
ck(_cdg._NOT_CHECKED and _cdg._NOT_CHECKED[0][0] == "talk time"
   and "no budget was recorded" in _cdg._NOT_CHECKED[0][1],
   "🔴 the Codex ledger RECORDS, behaviourally — asserting only that the function exists in the "
   "source passes with its body emptied, which a mutant proved")

print("\n— the audit: state, failure paths, and the runtimes that read a different file")
rd._NOT_CHECKED.clear(); rd._SECTIONS_RUN.clear()
rd._NOT_CHECKED["stale"] = "left over from an earlier pass"
rd._SECTIONS_RUN.extend(["stale", "older"])
_src_rd = (SKILL / "scripts" / "render_deck.py").read_text(encoding="utf-8")
ck("_NOT_CHECKED.clear()" in _src_rd and "_SECTIONS_RUN.clear()" in _src_rd,
   "🔴 a gate pass RESETS the ledger. It is module state, and a second pass in the same process — "
   "a test harness, an agent importing this module — inherited the first pass's rows and counted "
   "its sections again, so one tally described two runs")
rd._NOT_CHECKED.clear(); rd._SECTIONS_RUN.clear()
rd._COLLECTED = []
with rd._gate_section("failing"):
    rd.die("this deck is wrong in a specific way")
with rd._gate_section("passing"):
    pass
_checked, _rows = rd.coverage_ledger(rd._SECTIONS_RUN)
rd._COLLECTED = None
ck((_checked, _rows) == (2, []),
   "a gate that FAILED counts as BOUND, not as unchecked — it checked something and did not like "
   "it, which is the opposite of not having looked", (_checked, _rows))

if SOFFICE:
    surv = str(TMP / "survivor.pdf")
    import shutil as _sh
    _sh.copy(str(TMP / "a11y-deliverable.pdf"), surv)
    _before = os.path.getsize(surv)
    ck(rd.render_accessible_pdf("/nonexistent/soffice", _deck("x.pptx"), surv) is None
       and os.path.exists(surv) and os.path.getsize(surv) == _before,
       "🔴 a FAILED accessible export returns None and leaves the plain deliverable exactly where "
       "it was — it raised FileNotFoundError out of a public function until an audit called it "
       "with a soffice that does not exist")
    ghost = str(TMP / "ghost-notes.pdf")
    ck(rd.render_notes_pdf("/nonexistent/soffice", _deck("y.pptx"), ghost) is None
       and not os.path.exists(ghost),
       "...and a failed handout leaves no stray file behind")

_cx_doc = (SKILL / "references" / "codex-runtime.md").read_text(encoding="utf-8")
ck("-notes.pdf" in _cx_doc and "PDF/UA" in _cx_doc and "COVERAGE" in _cx_doc,
   "🔴 the CODEX runbook names both new deliverables and the coverage line. Documenting them only "
   "in the Claude-side files is how a runtime ships a deliverable nobody mentions — the same gap "
   "an earlier audit found for the Step-0 questions")
_ho = (SKILL / "references" / "handoff-checklist.md").read_text(encoding="utf-8")
ck("-notes.pdf" in _ho and "COVERAGE" in _ho,
   "...and the SHARED hand-off checklist requires naming them in the note, so the user hears about "
   "a file that was parked for them")

print()
print("%d passed, %d failed" % (len(OK), len(BAD)))
raise SystemExit(1 if BAD else 0)
