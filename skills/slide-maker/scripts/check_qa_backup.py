#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""A backup slide you cannot reach during the questions is one you did not prepare.

The practice this checks is ordinary in exactly the rooms this skill builds for — a defense, a
guidance committee, a conference talk, a board readout: you write down the questions you expect,
you build the slide that answers each, and you park them after the close. What makes the practice
work is not the slides. It is being able to JUMP to one while the room watches, and to get back.

Before `deckkit.link` / `agenda(targets=…)` / `back_link` existed, this library could not make a
deck jump at all (`grep hlinkClick` found it only in `_EA_FOLLOWERS`, an XML element-order constant), so a prepared
answer could only be reached by arrowing past every slide in between. This gate closes the loop the
way the rest of the skill does: the plan records the questions, and the BUILT file is checked
against them.

  "content": {"qa": [{"question": "Why not compressed sensing as the baseline?", "slide": 14},
                     {"question": "What is the failure mode at 12x?",            "slide": 15}]}

WHAT IT REFUSES TO DO. No `content.qa` recorded -> NOT CHECKED. Anticipating questions is a
practice, not a law, and a gate that demanded it of every deck would be teaching people to write
`[]` to make a red light go away.

    python3 scripts/check_qa_backup.py <deck.pptx> --gates <record.json> [--json] [--waive "<why>"]
    python3 scripts/check_qa_backup.py --selftest

Exit 0 clean · 1 findings · 2 could not run (NOT the same as clean, and it says so).
"""
from __future__ import annotations

import argparse
import json


def recorded_qa(gates):
    """The anticipated questions from either runtime's record, or None when none were recorded."""
    if not isinstance(gates, dict):
        return None
    for holder in ("content", "design_plan", "design"):
        blk = gates.get(holder)
        if isinstance(blk, dict) and isinstance(blk.get("qa"), list) and blk["qa"]:
            return blk["qa"]
    return None


def jump_map(prs):
    """{slide index (1-based): set of slide indices it JUMPS to} — real PowerPoint slide actions."""
    slides = list(prs.slides)
    # Keyed by the file's own slide IDs, not by Python identity. python-pptx happens to hand back
    # the same Slide object for a part twice, so `id()` works today — but that is a caching detail
    # of the library, and a jump map that silently comes back EMPTY if it ever changed would read
    # exactly like a deck with no links, which is the verdict this gate is built to distinguish.
    index = {}
    for i, sl in enumerate(slides, 1):
        index[id(sl)] = i
        sid = getattr(sl, "slide_id", None)
        if sid is not None:
            index[("id", sid)] = i
    out = {}
    for i, slide in enumerate(slides, 1):
        for shape in slide.shapes:
            try:
                target = shape.click_action.target_slide
            except Exception:
                target = None
            if target is None:
                continue
            at = index.get(("id", getattr(target, "slide_id", None)), index.get(id(target)))
            if at:
                out.setdefault(i, set()).add(at)
    return out


def check(pptx, qa, *, waive=None):
    """(findings, facts). findings = [(severity, code, message)] — 'block' | 'note'."""
    if not qa:
        raise RuntimeError("no anticipated questions recorded (`content.qa`) — nothing to check. "
                           "Preparing backup slides is a practice, not a law: a gate that demanded "
                           "it of every deck would just teach people to record an empty list")
    try:
        from pptx import Presentation
        prs = Presentation(pptx)
    except Exception as exc:
        raise RuntimeError("could not open %s: %s" % (pptx, exc))
    n = len(prs.slides)
    if not n:
        raise RuntimeError("the deck has no slides")
    jumps = jump_map(prs)
    reachable = {t for targets in jumps.values() for t in targets}
    findings, facts = [], {"slides": n, "questions": len(qa), "jumps": {k: sorted(v) for k, v in jumps.items()},
                           "unreachable": [], "no_way_back": []}

    for i, entry in enumerate(qa, 1):
        if not isinstance(entry, dict):
            findings.append(("block", "MALFORMED", "question %d is %r, not {\"question\": …, "
                                                   "\"slide\": …}" % (i, entry)))
            continue
        q = str(entry.get("question") or "").strip()
        slide = entry.get("slide")
        if not q:
            findings.append(("block", "NO QUESTION", "entry %d records a slide but not the question "
                                                     "it answers — the question is the part that "
                                                     "gets rehearsed" % i))
        if not isinstance(slide, int) or not 1 <= slide <= n:
            findings.append(("block", "NO SLIDE",
                             "%r points at slide %r, which this %d-slide deck does not have"
                             % (q[:60] or ("entry %d" % i), slide, n)))
            continue
        if slide not in reachable:
            facts["unreachable"].append(slide)
            findings.append(("block", "PREPARED BUT UNREACHABLE",
                             "slide %d answers %r and NOTHING links to it — during questions it can "
                             "only be reached by arrowing past every slide in between, which is why "
                             "prepared answers go unused. Link it: dk.link(shape, backup) from the "
                             "slide the question lands on, or an agenda row (targets=)."
                             % (slide, q[:70])))
        elif not jumps.get(slide):
            facts["no_way_back"].append(slide)
            findings.append(("note", "NO WAY BACK",
                             "slide %d can be jumped to but links nowhere — after the answer the "
                             "room watches you arrow backwards. dk.back_link(slide, home)." % slide))
    if waive and findings:
        facts["waived"] = waive
    return findings, facts


def _selftest():
    bad = []
    if recorded_qa({"content": {"qa": [{"question": "q", "slide": 3}]}}) != [{"question": "q", "slide": 3}]:
        bad.append("the shared record's qa list is not read")
    if recorded_qa({"design": {"qa": [{"question": "q", "slide": 3}]}}) is None:
        bad.append("the Codex twin is not read")
    if recorded_qa({"content": {"qa": []}}) is not None:
        bad.append("an empty list should read as 'nothing recorded', not as a filled one")
    if recorded_qa({}) is not None or recorded_qa(None) is not None:
        bad.append("an absent record should read as nothing")
    for b in bad:
        print("  ✗", b)
    print("[qa-backup] selftest %s" % ("FAILED" if bad else "ok"))
    return 1 if bad else 0


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0],
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx", nargs="?")
    ap.add_argument("--gates", help="a .deck-gates.json / codex evidence file holding content.qa")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--waive", default=None, help="a written reason; reports but does not fail")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args(argv)
    if a.selftest:
        return _selftest()
    if not a.pptx:
        ap.error("a deck path is required")
    qa = None
    if a.gates:
        try:
            qa = recorded_qa(json.load(open(a.gates, encoding="utf-8")))
        except Exception as exc:
            print("[qa-backup] NOT CHECKED — could not read %s: %s" % (a.gates, exc))
            return 2
    try:
        findings, facts = check(a.pptx, qa, waive=a.waive)
    except Exception as exc:
        print("[qa-backup] NOT CHECKED — %s" % exc)
        print("            NOT the same as clean.")
        return 2
    if a.json:
        print(json.dumps({"findings": [{"severity": s, "code": c, "why": m} for s, c, m in findings],
                          "facts": facts}, indent=1, ensure_ascii=False))
    print("[qa-backup] %d anticipated question(s) against a %d-slide deck"
          % (facts["questions"], facts["slides"]))
    for sev, code, why in findings:
        print("[qa-backup] %s %s: %s" % ("✗" if sev == "block" else "•", code, why))
    if facts.get("waived"):
        print("[qa-backup] WAIVED — %s" % facts["waived"])
        return 0
    if not findings:
        print("[qa-backup] every anticipated question has a backup slide, and every backup slide "
              "can be reached and left")
    return 1 if any(s == "block" for s, _c, _m in findings) else 0


if __name__ == "__main__":
    raise SystemExit(main())
