#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""The markers on the slides, the reference list, and the .bib must be the SAME three things.

An academic deck fails at citation in four ways, and all four are computable from the built file
plus the bibliography it claims to cite:

  * a marker on a slide that resolves to nothing — `[7]` over a six-entry list, the classic
    leftover from a cut slide, and the one an audience member WILL look up;
  * a reference list entry nothing cites — padding, and in a defense it reads as padding;
  * a reference list that was planned and never rendered, so the markers point at a page that is
    not in the deck;
  * an entry the bibliography cannot support — no author, no title, no year — which is where a
    retyped citation quietly becomes a misattribution.

The record names the bibliography and the keys, in the order they are cited:

  "content": {"citations": {"bib": "refs.bib", "style": "numeric",
                            "keys": ["lustig2007", "schlemper2018"]}}

`scripts/citations.py` derives both the marker and the reference line from the SAME entry, so this
gate never compares a slide against a second hand-typed copy — the failure it exists to catch.

WHAT IT REFUSES TO DO. No `content.citations` -> NOT CHECKED. Most decks cite nothing, and a gate
that demanded a bibliography of every deck would teach authors to record an empty one.

    python3 scripts/check_citations.py <deck.pptx> --gates <record.json> [--json] [--waive "<why>"]
    python3 scripts/check_citations.py --selftest

Exit 0 clean · 1 findings · 2 could not run (NOT the same as clean, and it says so).
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import citations as cit                                                   # noqa: E402

# 「（张三等，2020）」 — a Chinese deck writes its markers in full-width punctuation, and a scanner
# that only knows ASCII parentheses reports every one of them as uncited.
_WIDE = {"（": "(", "）": ")", "，": ", ", "［": "[", "］": "]", "、": ", ", "；": "; "}
_NUMERIC = re.compile(r"\[(\d+(?:\s*[,;–—-]\s*\d+)*)\]")
_AUTHOR_YEAR = re.compile(
    r"\(\s*([^()\d,;]{2,60}?)\s*,\s*(\d{4})[a-z]?\s*\)"        # (Lustig et al., 2007)
    r"|([A-Z一-鿿][^()\d,;]{1,40}?)\s*\(\s*(\d{4})[a-z]?\s*\)")   # Lustig et al. (2007)


def recorded_citations(gates):
    """The citation plan from either runtime's record, or None when none was recorded."""
    if not isinstance(gates, dict):
        return None
    for holder in ("content", "design_plan", "design"):
        blk = gates.get(holder)
        if isinstance(blk, dict) and isinstance(blk.get("citations"), dict):
            plan = blk["citations"]
            if plan.get("bib") and plan.get("keys"):
                return plan
    return None


def _widen(s):
    """Full-width punctuation -> ASCII, whitespace collapsed. CASE IS KEPT: the narrative marker
    `Zbontar (2018)` is recognised by its capital, so lowercasing first deletes half the matches."""
    for a, b in _WIDE.items():
        s = s.replace(a, b)
    return re.sub(r"\s+", " ", s).strip()


def _norm(s):
    return _widen(s).lower()


def slide_text(pptx):
    """[(slide text, notes text)] — 1-based by position, every shape's text flattened."""
    from pptx import Presentation
    out = []
    for slide in Presentation(pptx).slides:
        body = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                body.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    body.extend(c.text for c in row.cells)
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text
        except Exception:
            notes = ""
        out.append((" \n".join(body), notes))
    return out


def numeric_markers(text):
    """Every number cited by a `[3]` / `[3, 5]` / `[3-6]` marker in this text."""
    found = set()
    for group in _NUMERIC.findall(_widen(text).replace("–", "-").replace("—", "-")):
        parts = re.split(r"\s*[,;]\s*", group)
        for p in parts:
            span = re.match(r"^(\d+)\s*-\s*(\d+)$", p.strip())
            if span:
                lo, hi = int(span.group(1)), int(span.group(2))
                if lo <= hi and hi - lo < 200:
                    found.update(range(lo, hi + 1))
            elif p.strip().isdigit():
                found.add(int(p.strip()))
    return found


def author_year_markers(text):
    """{(first surname lowered, year)} for every `(Berg et al., 2020)` / `Berg (2020)` marker."""
    found = set()
    for a, b, c, d in _AUTHOR_YEAR.findall(_widen(text)):
        who, yr = (a, b) if a else (c, d)
        # `et al.` / `&` / `and` take spaces; 等 / 与 do not, because Chinese does not space them.
        who = re.split(r"(?:\s+(?:et al\.?|&|and)\s*|\s*[等与]\s*)", who.strip(), maxsplit=1)[0]
        who = who.strip(" .,;:、").lower()
        if who and not who.isdigit():
            found.add((who, yr))
    return found


def _rendered_on(pages, entry):
    """Which slide carries this entry's reference LINE, or None. Two tiers, so a hand-built list
    that does not match `format_reference` byte-for-byte is still recognised."""
    title = _norm(entry.get("title", ""))[:40]
    for i, (body, _notes) in enumerate(pages, 1):
        if title and title in _norm(body):
            return i
    names = cit.authors(entry)
    sn, yr = (_norm(cit.surname(names[0])) if names else ""), cit.year(entry)
    if sn and yr:
        for i, (body, _notes) in enumerate(pages, 1):
            nb = _norm(body)
            if sn in nb and yr in nb:
                return i
    return None


def check(pptx, plan, *, root=".", waive=None):
    """(findings, facts). findings = [(severity, code, message)] — 'block' | 'note'."""
    if not plan:
        raise RuntimeError("no citation plan recorded (`content.citations`) — nothing to check. "
                           "Most decks cite nothing, and a gate that demanded a bibliography of "
                           "every deck would only teach authors to record an empty one")
    style = plan.get("style") or "numeric"
    if style not in cit.STYLES:
        raise RuntimeError("unknown citation style %r — %s" % (style, " / ".join(cit.STYLES)))
    keys = list(plan.get("keys") or [])
    bib_rel = str(plan.get("bib") or "")
    # 🔴 The path comes out of a record a build wrote; it stays inside the deck folder.
    root_abs = os.path.realpath(root)
    bib = os.path.realpath(os.path.join(root_abs, bib_rel))
    if os.path.commonpath([bib, root_abs]) != root_abs:
        raise RuntimeError("the bibliography path %r resolves outside the deck folder — a record is "
                           "not a licence to read anywhere on the machine" % bib_rel)
    try:
        db = cit.parse_bibtex(open(bib, encoding="utf-8").read())
    except Exception as exc:
        raise RuntimeError("could not read the bibliography %s: %s" % (bib_rel, exc))
    try:
        pages = slide_text(pptx)
    except Exception as exc:
        raise RuntimeError("could not open %s: %s" % (pptx, exc))
    if not pages:
        raise RuntimeError("the deck has no slides")

    findings = []
    facts = {"slides": len(pages), "style": style, "entries": len(db), "cited": len(keys),
             "bib": bib_rel, "list_slides": [], "uncited": [], "dangling": []}

    entries, ok_keys = [], []
    for k in keys:
        e = db.get(k)
        if e is None:
            findings.append(("block", "NO SUCH ENTRY",
                             "%r is cited by the deck and is not in %s — a citation to a paper the "
                             "bibliography does not have cannot be checked, and cannot be looked "
                             "up by anyone in the room either" % (k, bib_rel)))
            continue
        gone = cit.missing(e)
        if gone:
            findings.append(("block", "INCOMPLETE ENTRY",
                             "%r has no %s. The line cannot be built without inventing the field, "
                             "and an invented citation is a claim about a real person's work — fill "
                             "it in the .bib (the DOI record has it) or drop the citation."
                             % (k, " or no ".join(gone))))
            continue
        entries.append(e)
        ok_keys.append(k)

    list_slides = set()
    for i, e in enumerate(entries, 1):
        on = _rendered_on(pages, e)
        if on:
            list_slides.add(on)
    facts["list_slides"] = sorted(list_slides)

    for i, (e, k) in enumerate(zip(entries, ok_keys), 1):
        if _rendered_on(pages, e) is None:
            findings.append(("block", "NOT IN THE LIST",
                             "%r is cited and appears in no reference list on any slide — the "
                             "marker points at a page that is not in the deck. Render it with "
                             "citations.reference_page(slide, entries)." % k))
            continue
        # cited = the marker appears somewhere that is NOT the reference list itself
        body_elsewhere = " \n".join(b for j, (b, _n) in enumerate(pages, 1) if j not in list_slides)
        notes_all = " \n".join(n for _b, n in pages)
        if style == "numeric":
            hit = i in numeric_markers(body_elsewhere) or i in numeric_markers(notes_all)
        else:
            names = cit.authors(e)
            want = (_norm(cit.surname(names[0])), cit.year(e))
            hit = want in author_year_markers(body_elsewhere) or want in author_year_markers(notes_all)
        if not hit:
            facts["uncited"].append(k)
            findings.append(("note", "UNCITED",
                             "%r is in the reference list and %s appears on no slide — an entry "
                             "nothing points to is padding, and in a defense it reads as padding."
                             % (k, cit.in_text(e, style, i))))

    body_elsewhere = " \n".join(b for j, (b, _n) in enumerate(pages, 1) if j not in list_slides)
    if style == "numeric":
        for n in sorted(numeric_markers(body_elsewhere)):
            if not 1 <= n <= len(entries):
                facts["dangling"].append(n)
                findings.append(("block", "DANGLING MARKER",
                                 "[%d] is cited on a slide and the reference list has %d entr%s — "
                                 "usually the leftover of a cut slide, and it is exactly what an "
                                 "audience member looks up. (On a numeric-style deck a bracketed "
                                 "INTEGER range reads as a citation too — `[12, 34] mm` is "
                                 "genuinely ambiguous, while `[0.2, 0.4]` is not. If that is what "
                                 "this is, waive it in writing.)"
                                 % (n, len(entries), "y" if len(entries) == 1 else "ies")))
    else:
        known = {(_norm(cit.surname(cit.authors(e)[0])), cit.year(e)) for e in entries
                 if cit.authors(e)}
        for who, yr in sorted(author_year_markers(body_elsewhere)):
            if (who, yr) not in known:
                facts["dangling"].append("%s %s" % (who, yr))
                findings.append(("block", "DANGLING MARKER",
                                 "(%s, %s) is cited on a slide and is in no reference list entry — "
                                 "a marker an audience member cannot resolve" % (who, yr)))
    if waive and findings:
        facts["waived"] = waive
    return findings, facts


def _selftest():
    bad = []
    if numeric_markers("as shown in [3] and [5, 7] and [10-12]") != {3, 5, 7, 10, 11, 12}:
        bad.append("numeric markers: %r" % numeric_markers("as shown in [3] and [5, 7] and [10-12]"))
    if numeric_markers("the 2020 cohort (n=14)") != set():
        bad.append("plain numbers must not read as citations")
    got = author_year_markers("as (Lustig et al., 2007) and Zbontar (2018) showed")
    if got != {("lustig", "2007"), ("zbontar", "2018")}:
        bad.append("author-year markers: %r" % got)
    if author_year_markers("（张三等，2020）") != {("张三", "2020")}:
        bad.append("a full-width CJK marker is not read: %r" % author_year_markers("（张三等，2020）"))
    if recorded_citations({"content": {"citations": {"bib": "r.bib", "keys": ["a"]}}}) is None:
        bad.append("the shared record's citation plan is not read")
    if recorded_citations({"content": {"citations": {"bib": "r.bib", "keys": []}}}) is not None:
        bad.append("an empty key list should read as 'nothing recorded'")
    if recorded_citations({}) is not None:
        bad.append("an absent record should read as nothing")
    for b in bad:
        print("  ✗", b)
    print("[citations-gate] selftest %s" % ("FAILED" if bad else "ok"))
    return 1 if bad else 0


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0],
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx", nargs="?")
    ap.add_argument("--gates", help="a .deck-gates.json / codex evidence file holding content.citations")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--waive", default=None, help="a written reason; reports but does not fail")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args(argv)
    if a.selftest:
        return _selftest()
    if not a.pptx:
        ap.error("a deck path is required")
    plan, root = None, os.path.dirname(os.path.abspath(a.pptx)) or "."
    if a.gates:
        try:
            plan = recorded_citations(json.load(open(a.gates, encoding="utf-8")))
            root = os.path.dirname(os.path.abspath(a.gates)) or "."
        except Exception as exc:
            print("[citations-gate] NOT CHECKED — could not read %s: %s" % (a.gates, exc))
            return 2
    try:
        findings, facts = check(a.pptx, plan, root=root, waive=a.waive)
    except Exception as exc:
        print("[citations-gate] NOT CHECKED — %s" % exc)
        print("                 NOT the same as clean.")
        return 2
    if a.json:
        print(json.dumps({"findings": [{"severity": s, "code": c, "why": m} for s, c, m in findings],
                          "facts": facts}, indent=1, ensure_ascii=False))
    print("[citations-gate] %d cited key(s), %s style, reference list on slide(s) %s"
          % (facts["cited"], facts["style"], facts["list_slides"] or "—"))
    for sev, code, why in findings:
        print("[citations-gate] %s %s: %s" % ("✗" if sev == "block" else "•", code, why))
    if facts.get("waived"):
        print("[citations-gate] WAIVED — %s" % facts["waived"])
        return 0
    if not findings:
        print("[citations-gate] every marker resolves, every entry is cited, and every line comes "
              "from the .bib")
    return 1 if any(s == "block" for s, _c, _m in findings) else 0


if __name__ == "__main__":
    raise SystemExit(main())
